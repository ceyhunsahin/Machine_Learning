# -*- coding: utf-8 -*-
import sys
import os
import collections
import base64
import datetime
import time
import json
import io
import dash
import dash_bootstrap_components as dbc
import dash_core_components as dcc
import dash_daq as daq  # #
import dash_html_components as html
import dash_table  # #
import pandas as pd
import plotly.graph_objects as go
from dash import no_update
from dash.dependencies import Input, Output, State
from dash.exceptions import PreventUpdate
from urllib.parse import quote as urlquote
from numpy import trapz
from flask import send_file
from openpyxl import Workbook, load_workbook


# from sshtunnel import SSHTunnelForwarder
# import mariadb
# import pywintypes
# pywintypes.datetime = pywintypes.TimeType

def find_data_file(filename):
    if getattr(sys, 'frozen', False):
        # The application is frozen
        datadir = os.path.dirname(sys.executable)
    else:
        # The application is not frozen
        # Change this bit to match where you store your data files:
        datadir = os.path.dirname(__file__)

    return os.path.join(datadir, filename)


# Initialize the app
app = dash.Dash(__name__, external_stylesheets=[dbc.themes.BOOTSTRAP], assets_folder=find_data_file('assets/'))
server = app.server
app.config.suppress_callback_exceptions = True

# connect OPC

# get data from MAF

getDataFromModbus = []

extra_data_list = [
    'Mass de Bois', 'Volume gaz', 'Vitesse de rotation', 'Puissance Thermique',
    'Puissance Electrique', 'CO', 'CO2', 'NO', 'NOX', 'Temperature de Fumée'
]

app.layout = html.Div([
    dcc.Location(id='url', refresh=False),
    html.Div(id='page-content'),
])
# 4 page
index_page = html.Div(className="indexpage",
                      children=[
                          dcc.Link(html.Button('Go to ENERBAT', id="indexPageStyle"), href='/page-1'),
                          html.Br(),
                          dcc.Link(html.Button('Go to X', id="indexPageStyle"), href='/page-2'),
                          html.Br(),
                          dcc.Link(html.Button('Go to Y', id="indexPageStyle"), href='/page-3'),
                          html.Br(),
                          dcc.Link(html.Button('Go to Z', id="indexPageStyle"), href='/page-4'),
                      ])

page_1_layout = html.Div(
    className='main_container',
    children=[
        html.Div(id='fourcolumnsdivusercontrols', className="four-columns-div-user-controls",
                 children=[
                     html.Div([daq.PowerButton(id='my-toggle-switch',
                                               label={'label': 'Connect OPC',
                                                      'style': {'fontSize': '22px', 'fontWeight': "bold"}},
                                               labelPosition='bottom', on=False, size=100, color="green",
                                               className='dark-theme-control'), html.Div(
                         dcc.Upload(
                             id='upload-data',
                             children=html.Div([
                                 'Drag and Drop or ',
                                 html.A('Select Files for work')
                             ]),
                             style={
                                 'visibility': 'hidden',
                             },
                             # Allow multiple files to be uploaded
                             multiple=True,

                         ),

                     )], ),

                     html.Div(id="openOPCDiv", children=[], style={'visibility': 'hidden'}),
                     html.Div(className='userControlDownSide',
                              children=[
                                  html.Div(className='userControlDownLeftSide',
                                           children=[
                                               html.Div(id="opcLoad",
                                                        className='div-for-dropdown',
                                                        children=[], ),
                                               html.Div(dcc.Interval(
                                                   id='interval',
                                                   interval=5000,
                                                   n_intervals=3,

                                               )),
                                               # html.Div(id = 'ceyhun',
                                               #          style = {'visibility' : 'hidden', 'height' :'1rem' }),
                                               html.Div(className="file_db_button",
                                                        children=[
                                                            html.Button('File', id='file_save', n_clicks=0, ),
                                                            html.Button('Database', id='db_save', n_clicks=0, ),
                                                        ]),
                                               html.Div(dcc.Store(id='datastore')),
                                               html.Div(id='pointLeftFirst', children=[], style={'display': 'None'}),
                                               html.Div(id='pointLeftSecond', children=[], style={'display': 'None'}),
                                               html.Div(id='pointRightFirst', children=[], style={'display': 'None'}),
                                               html.Div(id='pointRightSecond', children=[], style={'display': 'None'}),
                                               html.Div(id='pointLeftFirstTab4', children=[],
                                                        style={'display': 'None'}),
                                               html.Div(id='pointLeftSecondTab4', children=[],
                                                        style={'display': 'None'}),
                                               html.Div(id='pointRightFirstTab4', children=[],
                                                        style={'display': 'None'}),
                                               html.Div(id='pointRightSecondTab4', children=[],
                                                        style={'display': 'None'}),
                                               html.Div(id='leftSideChecklistValue', children=[],
                                                        style={'display': 'None'}),
                                               html.Div(id='leftSidedroptValue', children=[],
                                                        style={'display': 'None'}),
                                               html.Div(id='leftSideChecklistValueHidden', children=[],
                                                        style={'display': 'None'}),
                                               html.Div(id='deletedval', children=[], style={'display': 'None'}),
                                               html.Div(id='leftSideChecklistValueHiddenTab4', children=[],
                                                        style={'display': 'None'}),
                                               html.Div(id='tab2hiddenValuex_axis', style={'display': 'None'},
                                                        children=[]),
                                               html.Div(id='tab2hiddenValuey_axis', style={'display': 'None'},
                                                        children=[]),
                                               html.Div(id='tab4hiddenValuex_axissecond', style={'display': 'None'},
                                                        children=[]),
                                               html.Div(id='tab4hiddenValuey_axissecond', style={'display': 'None'},
                                                        children=[]),
                                               html.Div(id='tab4hiddenValuex_axis', style={'display': 'None'},
                                                        children=[]),

                                               html.Div(id='tab3hiddenValuey_axis', style={'display': 'None'},
                                                        children=[]),
                                               html.Div(id='hiddenTextHeader', children=[], style={'display': 'None'}),
                                               html.Div(id='hiddenTextNote', children=[], style={'display': 'None'}),
                                               html.Div(id='hiddenTextxaxis', children=[], style={'display': 'None'}),
                                               html.Div(id='hiddenTextyaxis', children=[], style={'display': 'None'}),
                                               html.Div(id='hiddenTextHeader4', children=[], style={'display': 'None'}),
                                               html.Div(id='hiddenTextNote4', children=[], style={'display': 'None'}),
                                               html.Div(id='hiddenTextxaxis4', children=[], style={'display': 'None'}),
                                               html.Div(id='hiddenTextyaxis4', children=[], style={'display': 'None'}),
                                               html.Div(id='hiddenShapeVal', children=[], style={'display': 'None'}),
                                               html.Div(id='hiddenShapeDate', children=[],
                                                        style={'display': 'None'}), ], ),
                                  html.Div(id='hiddenDifferance', children=[], style={'display': 'None'}),
                                  html.Div(id='retrieve', children=[], style={'display': 'None'}),
                                  html.Div(id='datatablehidden', children=[], style={'display': 'None'}),
                                  html.Div(id='radiographhidden', children=[], style={'display': 'None'}),
                                  html.Div(id='sliderHeightTab1hidden', children=[], style={'display': 'None'}),
                                  html.Div(id='sliderWidthTab1hidden', children=[], style={'display': 'None'}),
                                  html.Div(id='hiddenShapeValtab4', children=[], style={'display': 'None'}),
                                  html.Div(id='hiddenShapeDatetab4', children=[], style={'display': 'None'}),
                                  html.Div(id='hiddenDifferancetab4', children=[], style={'display': 'None'}),
                                  html.Div(id='retrievetab4', children=[], style={'display': 'None'}),
                                  html.Div(id='datatablehiddentab4', children=[], style={'display': 'None'}),
                                  html.Div(id='radiographhiddentab4', children=[], style={'display': 'None'}),
                                  html.Div(id='sliderHeightTab1hiddentab4', children=[], style={'display': 'None'}),
                                  html.Div(id='sliderWidthTab1hiddentab4', children=[], style={'display': 'None'}),
                                  html.Div(id='minimumValueGraphhiddenfirst', children=[], style={'display': 'None'}),
                                  html.Div(id='minimumValueGraphhiddensecond', children=[], style={'display': 'None'}),
                                  html.Div(id='firstchoosenvalhidden', children=[], style={'display': 'None'}),
                                  html.Div(id='secondchoosenvalhidden', children=[], style={'display': 'None'}),
                                  html.Div(id='leftintegralfirsthidden', children=[], style={'display': 'None'}),
                                  html.Div(id='leftintegralsecondhidden', children=[], style={'display': 'None'}),
                                  html.Div(id='rightintegralfirsthidden', children=[], style={'display': 'None'}),
                                  html.Div(id='rightintegralsecondhidden', children=[], style={'display': 'None'}),
                                  html.Div(id='tableinteractivehidden', children=[], style={'display': 'None'}),
                                  html.Div(id='firstchoosenvalhiddentab4', children=[], style={'display': 'None'}),
                                  html.Div(id='secondchoosenvalhiddentab4', children=[], style={'display': 'None'}),
                                  html.Div(id='leftintegralfirsthiddentab4', children=[], style={'display': 'None'}),
                                  html.Div(id='leftintegralsecondhiddentab4', children=[], style={'display': 'None'}),
                                  html.Div(id='rightintegralfirsthiddentab4', children=[], style={'display': 'None'}),
                                  html.Div(id='rightintegralsecondhiddentab4', children=[], style={'display': 'None'}),
                                  html.Div(id='tableinteractivehiddentab4', children=[], style={'display': 'None'}),
                                  html.Div(id='writeexcelhidden', children=[], style={'display': 'None'}),
                                  html.Div(id='writeexcelhiddenTab4', children=[], style={'display': 'None'}),
                                  html.Div(id='hiddenrecord1', children=[], style={'display': 'None'}),
                                  html.Div(id='hiddenrecord2', children=[], style={'display': 'None'}),
                                  html.Div(id='hiddenrecord3', children=[], style={'display': 'None'}),
                                  html.Div(id='hiddenrecord4', children=[], style={'display': 'None'}),
                                  html.Div(id='inputRightY_axishidden', children=[], style={'display': 'None'}),
                                  html.Div(id='inputRightX_axishidden', children=[], style={'display': 'None'}),
                                  html.Div(id='valueSendRighthidden', children=[], style={'display': 'None'}),
                                  html.Div(id='checklistvaleurhidden', children=[], style={'display': 'None'}),
                                  html.Div(id='shiftaxisdrophidden', children=[], style={'display': 'None'}),
                                  html.Div(id='shift_x_axishidden', children=[], style={'display': 'None'}),
                                  html.Div(id='shift_y_axishidden', children=[], style={'display': 'None'}),
                                  html.Div(id='tab1sendhidden', children=[], style={'display': 'None'}),
                                  html.Div(id='shiftaxisdroptab4hidden', children=[], style={'display': 'None'}),
                                  html.Div(id='shift_x_axistab4hidden', children=[], style={'display': 'None'}),
                                  html.Div(id='shift_y_axistab4hidden', children=[], style={'display': 'None'}),
                                  html.Div(id='output_s', children=[], style={'display': 'None'}),
                                  html.Div(id='radiographtab4hidden', children=[], style={'display': 'None'}),
                                  html.Div(dcc.Graph(id='graphhidden',
                                                     config={},
                                                     style={'display': 'None'},
                                                     figure={
                                                         'layout': {'legend': {'tracegroupgap': 0},

                                                                    }
                                                     }

                                                     ), ),
                                  html.Div(dcc.Graph(id='graphTab4hidden',
                                                     config={},
                                                     style={'display': 'None'},
                                                     figure={
                                                         'layout': {'legend': {'tracegroupgap': 0},

                                                                    }
                                                     }

                                                     ), )
                                  ,

                              ]),
                 ]),

        html.Div(id='eightcolumnsdivforcharts', className="eight-columns-div-for-charts",
                 children=[
                     html.Div(
                         className='right-upper',
                         children=[
                             html.Div([
                                 dcc.Tabs(
                                     id="tabs-with-classes",
                                     value='tab-1',
                                     parent_className='custom-tabs',
                                     className='custom-tabs-container',
                                     children=[
                                         dcc.Tab(
                                             id="tab1",
                                             label='Work on unique parameter',
                                             value='tab-1',
                                             className='custom-tab',
                                             selected_className='custom-tab--selected',
                                             children=[],
                                         ),
                                         dcc.Tab(
                                             id='tab2',
                                             label='Tab for one option',
                                             value='tab-2',
                                             className='custom-tab',
                                             selected_className='custom-tab--selected',
                                             children=[
                                             ]
                                         ),
                                         dcc.Tab(
                                             id='tab3',
                                             label='Work On Database',
                                             value='tab-3', className='custom-tab',

                                             # style = {'visibility' : 'hidden'},
                                             selected_className='custom-tab--selected'
                                         ),
                                         dcc.Tab(
                                             id="tab4",
                                             label='Work on different parameters',
                                             value='tab-4',
                                             className='custom-tab',
                                             # style={'visibility': 'hidden'},
                                             selected_className='custom-tab--selected'
                                         ),
                                         dcc.Tab(
                                             id="tab5",
                                             label='Tab for one option',
                                             value='tab-5',
                                             className='custom-tab',
                                             style={'visibility': 'hidden'},
                                             selected_className='custom-tab--selected'
                                         ),
                                     ]),
                                 html.Div(id='tabs-content-classes'),

                             ]),

                         ]),

                 ]
                 ),
        # dcc.Graph(id = "first_value_graph", config = {'displayModeLine': True}, animate=True)
    ]),

page_2_layout = html.Div([
    dcc.Link('Go to MODBUS', href='/page-1'),
    html.Br(),
    dcc.Link('Go to Y', href='/page-3'),
    html.Br(),
    dcc.Link('Go to Z', href='/page-4'),
    html.Br(),
    dcc.Link('Go back to home', href='/')
])

page_3_layout = html.Div([
    dcc.Link('Go to MODBUS', href='/page-1'),
    html.Br(),
    dcc.Link('Go to X', href='/page-2'),
    html.Br(),
    dcc.Link('Go to Z', href='/page-4'),
    html.Br(),
    dcc.Link('Go back to home', href='/')
])

page_4_layout = html.Div([
    dcc.Link('Go to MODBUS', href='/page-1'),
    html.Br(),
    dcc.Link('Go to X', href='/page-2'),
    html.Br(),
    dcc.Link('Go to Y', href='/page-3'),
    html.Br(),
    dcc.Link('Go back to home', href='/')
])


# @app.callback(Output('tab2', 'children'),
#               [Input("my-toggle-switch", "on"), Input('interval', 'n_intervals')])
# def values(on, n_intervals):
#     if on == 1:
#
#         opc = OpenOPC.client()
#         opc.servers()
#         opc.connect('Kepware.KEPServerEX.V6')
#
#         for name, value, quality, time in opc.iread(
#                 ['schneider_Xflow.MAF.CoAd', 'schneider_Xflow.MAF.ComManCoP2',
#                  'schneider_Xflow.MAF.ComManCoP3P4P5', 'schneider_Xflow.MAF.ComManPompeSec',
#                  'schneider_Xflow.MAF.CompteurEnergie', 'schneider_Xflow.MAF.CoP2',
#                  'schneider_Xflow.MAF.CtempDepChauff', 'schneider_Xflow.MAF.D1',
#                  'schneider_Xflow.MAF.D2', 'schneider_Xflow.MAF.D3', 'schneider_Xflow.MAF.D4',
#                  'schneider_Xflow.MAF.MarcheBruleur', 'schneider_Xflow.MAF.Teg',
#                  'schneider_Xflow.MAF.SdeBasBouMelange', 'schneider_Xflow.MAF.SdeBasHauMelange',
#                  'schneider_Xflow.MAF.TambN3', 'schneider_Xflow.MAF.Tb1', 'schneider_Xflow.MAF.Tb2',
#                  'schneider_Xflow.MAF.Tb3', 'schneider_Xflow.MAF.Tb4', 'schneider_Xflow.MAF.TdepPLC',
#                  'schneider_Xflow.MAF.Teb', 'schneider_Xflow.MAF.Tec ', 'schneider_Xflow.MAF.Teev',
#                  'schneider_Xflow.MAF.TempminMaf', 'schneider_Xflow.MAF.Text', 'schneider_Xflow.MAF.Tsb',
#                  'schneider_Xflow.MAF.Tsc', 'schneider_Xflow.MAF.Tsev', 'schneider_Xflow.MAF.Tsg']):
#             getDataFromModbus.append((name, value, quality, time))
#             df = pd.DataFrame(getDataFromModbus, columns=['ItemID', 'Value', 'DataType', 'TimeStamp'])
#             df.to_csv("cc.csv")
#     return getDataFromModbus

# surf between pages
# Update the index
@app.callback(dash.dependencies.Output('page-content', 'children'),
              [dash.dependencies.Input('url', 'pathname')])
def display_page(pathname):
    if pathname == '/page-1':
        return page_1_layout
    elif pathname == '/page-2':
        return page_2_layout
    elif pathname == '/page-3':
        return page_3_layout
    elif pathname == '/page-4':
        return page_4_layout
    else:
        return index_page


def parse_contents(contents, filename, date):
    content_type, content_string = contents.split(',')

    decoded = base64.b64decode(content_string)
    try:
        if 'csv' in filename:
            # Assume that the user uploaded a CSV file
            df = pd.read_csv(
                io.StringIO(decoded.decode('utf-8')))
        elif 'xlsx' in filename:
            # Assume that the user uploaded an excel file
            df = pd.read_excel(io.BytesIO(decoded))
            df['index'] = df.index
            df = df.reindex(columns=sorted(df.columns, reverse=True))
            df.to_excel("appending.xlsx")
            df.to_excel("rawinfo.xlsx")
    except Exception as e:
        print(e)
        return html.Div([
            'There was an error processing this file.'
        ])

    return html.Div([
        html.H5(filename),
        html.H6(datetime.datetime.fromtimestamp(date)),
        dash_table.DataTable(
            id='datatable-interactivity',
            data=df.to_dict('records'),
            columns=[{'name': i, 'id': i, "deletable": True, "selectable": True} for i in df.columns if
                     i[1:].isdigit() != 1],
            editable=True,
            page_size=50,
            style_table={'height': '500px', 'overflowY': 'auto', 'width': '98%'},
            style_cell={
                'overflow': 'hidden',
                'textOverflow': 'ellipsis',
                'maxWidth': 0,
                'fontSize': '1rem',
                'TextAlign': 'center',
            },
            fixed_rows={'headers': True},
            tooltip_data=[
                {
                    column: {'value': str(value), 'type': 'markdown'}
                    for column, value in row.items()
                } for row in df.to_dict('records')
            ],
            style_cell_conditional=[
                {
                    'if': {'column_id': c},
                    'textAlign': 'center',
                    'width': '8%'}

                for c in df.columns if c != 'date'],
            # style_cell_conditional=[
            # {'if': {'column_id': 'date'},
            #  'width': '15%'}

            style_header={
                'backgroundColor': 'rgb(230, 230, 230)',
                'fontWeight': 'bold'
            },
            filter_action="native",
            sort_action="native",
            sort_mode="multi",
            column_selectable="single",
            # row_selectable="multi",
            # row_deletable=True,
            selected_columns=[],
            selected_rows=[],
            page_action="native",
            page_current=0,
            export_format='xlsx',
            export_headers='display',
            merge_duplicate_headers=True
        ),

        html.Hr(),  # horizontal line
    ])


@app.callback([Output('datatablehidden', 'children'), Output('retrieve', 'children')],
              [Input('upload-data', 'contents'), Input("my-toggle-switch", "on"), ],
              [State('upload-data', 'filename'),
               State('upload-data', 'last_modified'),
               State('retrieve', 'children'),
               State('datatablehidden', 'children')])
def update_output(list_of_contents, on, list_of_names, list_of_dates, retrieve, content):
    if on == 0:
        raise PreventUpdate
    if list_of_contents is not None:

        content = [
            parse_contents(c, n, d) for c, n, d in
            zip(list_of_contents, list_of_names, list_of_dates)]
        retrieve = list_of_names

        return content, retrieve
    else:
        return (no_update, no_update)


@app.callback(Output('output-data-upload', 'children'),
              [Input('datatablehidden', 'children')],
              )
def retrieve(retrieve):
    if retrieve == None:
        raise PreventUpdate
    return retrieve


# @app.callback(Output('tab2DashTable', 'children'),
#               [Input('datatablehidden', 'children')],
#               )
# def retrieve2(retrieve):
#     return retrieve

@app.callback(Output('tab4DashTable', 'children'),
              [Input('datatablehidden', 'children')],
              )
def retrieve4(retrieve):
    return retrieve


@app.callback(
    Output('datatable-interactivity', 'style_data_conditional'),
    [Input('datatable-interactivity', 'selected_columns')]
)
def update_styles(selected_columns):
    return [{
        'if': {'column_id': i},
        'background_color': 'red'
    } for i in selected_columns]

    # Output("opcLoad","children") : for load left and right side,
    # for this created a hiddev div as opcLoad,
    # Output('tab2','children') : also hidden tab, for the graph


@app.callback([Output("opcLoad", "children"), Output('upload-data', 'style')],
              [Input("my-toggle-switch", "on")]
              )
def opcLoadingData(on):
    ocploadlist = []

    if on == 1:
        visibilty = {'width': '100%',
                     'height': '35px',
                     'lineHeight': '25px',
                     'borderWidth': '1px',
                     'borderStyle': 'dashed',
                     'borderRadius': '5px',
                     'textAlign': 'center',
                     'margin': '20px',
                     'visibility': 'visible'}
        data_list = ['CoAd', 'ComManCoP2', 'ComManCoP3P4P5', 'ComManPompeSec', 'CompteurEnergie', 'CoP2',
                     'CtempDepChauff',
                     'D1', 'D2', 'D3', 'D4', 'MarcheBruleur', 'Teg', 'SdeBasBouMelange', 'SdeBasHauMelange', 'TambN3',
                     'Tb1',
                     'Tb2', 'Tb3', 'Tb4', 'TdepPLC', 'Teb', 'Tec', 'Teev', 'TempminMaf', 'Text', 'Tsb', 'Tsc', 'Tsev']

        ocploadlist = html.Div(className="userControlDownSideCreated",
                               children=[html.Div(className="userControlDownLeftSide",

                                                  children=[html.Div(className='aa',
                                                                     children=[html.Div(
                                                                         dcc.Dropdown(id='dropdownLeft',
                                                                                      options=[{'label': i, 'value': i}
                                                                                               for i in data_list if
                                                                                               i != 'date'],
                                                                                      multi=False,
                                                                                      style={"cursor": "pointer"},
                                                                                      className='stockSelectorClass',
                                                                                      clearable=False,
                                                                                      placeholder='Select your parameters...',
                                                                                      ),
                                                                     ),
                                                                         html.Div([html.Button('Show', id='showLeft',
                                                                                               n_clicks=0,
                                                                                               style={'height': '40px',
                                                                                                      'width': '80px',
                                                                                                      'fontSize': '1.2rem'}),
                                                                                   html.Button('Delete', id='clearLeft',
                                                                                               n_clicks=0,
                                                                                               style={'height': '40px',
                                                                                                      'width': '80px',
                                                                                                      'fontSize': '1.2rem'})],
                                                                                  className='buttons'),
                                                                         html.Div(id='leftSideDropdownHidden',
                                                                                  children=[],
                                                                                  style={'display': 'None'}),
                                                                         # html.Div(id='leftSideDropdown', children=[]),
                                                                         html.Div([dbc.Checklist(
                                                                             id='choosenChecklistLeft',
                                                                             options=[{'label': i, 'value': i} for i in
                                                                                      []],
                                                                             value=[],
                                                                             labelStyle={'display': 'Block'},
                                                                         ), ], style={"marginTop": "8px",
                                                                                      "marginLeft": "8px",
                                                                                      'visibility': 'hidden'}),
                                                                         html.Div(
                                                                             [

                                                                                 dbc.Modal(
                                                                                     [
                                                                                         dbc.ModalHeader("INFORMATION"),
                                                                                         dbc.ModalBody(
                                                                                             "Vous pouvez choisir maximum 20 valeur"),
                                                                                         dbc.ModalFooter(
                                                                                             dbc.Button("Close",
                                                                                                        id="close",
                                                                                                        className="ml-auto")
                                                                                         ),
                                                                                     ],
                                                                                     id="modal",
                                                                                 ),
                                                                             ])
                                                                     ])], ),
                                         html.Div(className="userControlDownRightSide",
                                                  children=[
                                                      html.Div(
                                                          className='div-for-dropdown',
                                                          children=[
                                                              html.Div(
                                                                  dcc.Dropdown(id='dropdownRight',
                                                                               options=[{'label': i, 'value': i} for i
                                                                                        in extra_data_list],
                                                                               multi=False,
                                                                               value='',
                                                                               style={"cursor": "pointer"},
                                                                               className='stockSelectorClass',
                                                                               clearable=False,
                                                                               placeholder='Select your parameters...',
                                                                               ),
                                                              ),
                                                              html.Div([html.Button('Show', id='showRight', n_clicks=0,
                                                                                    style={'height': '40px',
                                                                                           'width': '80px',
                                                                                           'fontSize': '1.2rem'}),
                                                                        html.Button('Delete', id='clearRight',
                                                                                    n_clicks=0, style={'height': '40px',
                                                                                                       'width': '80px',
                                                                                                       'fontSize': '1.2rem'})],
                                                                       className='buttons'),
                                                              html.Div(id='rightSideDropdownHidden', children=[],
                                                                       style={'visibility': 'hidden'}),
                                                              html.Div(id="rightSideDropdown", children=[])
                                                          ],
                                                      ),
                                                  ]),
                                         ])
        return (ocploadlist, visibilty)

    else:
        return (ocploadlist, {'visibility': 'hidden'})


@app.callback(Output("dropdownLeft", "options"),
              [Input("retrieve", "children")])
def dropdownlistcontrol(retrieve):
    if len(retrieve) > 0:
        df = pd.read_excel('appending.xlsx')
        dff = [{'label': i, 'value': i} for i in df.columns if i.startswith('Un') != 1 and i != 'index' and i != 'date']
        return dff
    else:
        return no_update


# @app.callback(
#     [Output("leftSideDropdownHidden", "children"),
#      Output("leftSidedroptValue", "children")],
#     [Input("dropdownLeft", "value"),],
#     [State("leftSideDropdownHidden", "children")]
# )
# def hiddendiv(val_dropdownLeft, children):
#     if val_dropdownLeft == None or val_dropdownLeft == '':
#         raise PreventUpdate
#     a = []
#     a.append(val_dropdownLeft)
#     for i in a:
#         if i not in children:
#             children.append(val_dropdownLeft)
#     return children, children
#

@app.callback(
    [Output('choosenChecklistLeft', 'options'),
     Output('choosenChecklistLeft', 'style'),
     Output('choosenChecklistLeft', 'value'),
     Output("leftSideDropdownHidden", "children"),
     Output("leftSidedroptValue", "children"),
     Output("deletedval", "children")],
    [Input("showLeft", "n_clicks"),
     Input("clearLeft", "n_clicks"),
     ],
    [State("dropdownLeft", "value"),
     State("leftSideDropdownHidden", "children"),
     State('choosenChecklistLeft', 'value'),
     State('deletedval', 'children')],
)
# left side dropdown-checklist relation
#########

def displayLeftDropdown(n_clicks1, nc2, dropval, valeur, value, deletedval):
    if dropval == None or deletedval == None:
        raise PreventUpdate
    q1 = dash.callback_context.triggered[0]["prop_id"].split(".")[0]
    a = []
    a.append(dropval)
    for i in a:
        if q1 == 'showLeft' and i not in valeur:
            valeur.append(dropval)
        if q1 == 'clearLeft' and i not in deletedval:
            pass
    if q1 == 'showLeft':
        return [{'label': i, 'value': i} for i in valeur], {'visibility': 'visible'}, [], valeur, valeur, deletedval

    if q1 == 'clearLeft':
        print('nclick ne oldu', nc2)

        for k in range(len(value)):
            valeur.remove(value[k])
            deletedval.append(value[k])

        return [{'label': i, 'value': i} for i in valeur], {'visibility': 'visible'}, [], valeur, valeur, deletedval
    else:
        no_update, no_update, no_update, no_update, no_update, no_update


@app.callback(
    Output("modal", "is_open"),
    [Input("showLeft", "n_clicks"), Input("close", "n_clicks")],
    [State("modal", "is_open"),
     State("leftSideDropdownHidden", "children")],
)
def toggle_modal(n1, n2, is_open, children):
    if len(children) > 20:
        return not is_open
    return is_open


#### rightside dropdown-checklist relation


@app.callback(
    [Output('rightSideDropdown', "children"), Output('checklistvaleurhidden', "children"), ],
    [
        Input("showRight", "n_clicks"),
        Input("clearRight", "n_clicks")
    ],
    [
        State("dropdownRight", "value"),
        State('rightSideDropdown', "children"),
        State('checklistvaleurhidden', "children")
    ]
)
def edit_list2(ncr1, ncr2, valeur, children, hiddenchild):
    triggered_buttons = dash.callback_context.triggered[0]["prop_id"].split(".")[0]

    if triggered_buttons == "showRight":
        def mesure1(textRight):
            if textRight == "Mass de Bois":
                return "g"
            elif textRight == 'Volume gaz':
                return 'm3'

            elif textRight == 'Vitesse de rotation':
                return 'tour/mn'

            elif textRight in {'Puissance Thermique', 'Puissance Electrique'}:
                return "W"

            elif textRight in {'CO', 'CO2', 'NO', 'NOX', 'CX'}:
                return "% MOL"


            elif textRight == 'Temperature de Fumée':
                return '°K'

        if valeur != '':
            new_listRight = html.Div([html.Div([
                html.Div([dcc.Markdown('''*{}'''.format(valeur), id="checklistValeur",
                                       style={'height': '1rem', 'fontFamily': 'arial', 'color': 'black',
                                              'fontSize': '1.2rem'}),
                          html.Div([dbc.Input(id='inputRightY_axis',
                                              type="text",
                                              min=-10000, max=10000, step=1, bs_size="sm", style={'width': '6rem'},
                                              placeholder='Y axis value',
                                              autoFocus=True, ),
                                    dbc.Input(id='inputRightX_axis',
                                              type="text",
                                              min=-10000, max=10000, step=1, bs_size="sm", style={'width': '6rem'},
                                              placeholder='X axis value',
                                              autoFocus=True, ),
                                    ], id="styled-numeric-input", ),
                          html.P(mesure1(valeur),
                                 style={'margin': '0.1rem 0', 'color': 'black', 'height': '2rem', 'fontFamily': 'arial',
                                        'fontSize': '1.2rem', }),
                          dbc.Button("Ok", id="valueSendRight", outline=True, n_clicks=0, color="primary",
                                     className="mr-1"),

                          ], className='design_children2'),
            ], className='design_children', ), html.Hr()])
            hiddenchild.append(valeur)

            children.append(new_listRight)

    if triggered_buttons == "clearRight":
        if len(children) == 0:
            raise PreventUpdate
        else:
            children.pop()

    return children, hiddenchild


@app.callback(Output('tabs-content-classes', 'children'),
              [Input('tabs-with-classes', 'value')],
              )
def render_content(tab):
    if tab == 'tab-1':
        return html.Div([
            html.Div(id='tab1Data')
        ])
    if tab == 'tab-2':
        return html.Div([
            html.Div(id='tab2Data')
        ])
    if tab == 'tab-3':
        return html.Div([
            html.Div(id='tab3Data', children=[]),
            html.Div(id='Dbdesign')])

    if tab == 'tab-4':
        return html.Div([
            html.Div(id='tab4Data')
        ])
    else:
        pass


@app.callback(Output('tab1Data', 'children'),
              [Input("my-toggle-switch", "on"),
               Input("leftSidedroptValue", "children")],
              [State('tabs-with-classes', 'value')]
              )
def LoadingDataTab1(on, dropdownhidden, tab):
    if on == 1 and tab == 'tab-1':
        loadTab1 = html.Div([html.Div([html.Div([html.Div([dcc.Dropdown(id='firstChoosenValue',
                                                                        options=[{'label': i, 'value': i} for i in
                                                                                 dropdownhidden],
                                                                        multi=False,
                                                                        style={"cursor": "pointer", 'width': '180px'},
                                                                        className='',
                                                                        clearable=True,
                                                                        placeholder='First Value...',
                                                                        ),
                                                           dbc.Input(id='leftIntegralFirst',
                                                                     type="text",
                                                                     debounce=True,
                                                                     min=-10000, max=10000, step=1,
                                                                     bs_size="sm",
                                                                     style={'width': '7rem', "marginTop": "1.5rem"},
                                                                     autoFocus=True,
                                                                     placeholder="first point"),
                                                           dbc.Input(id='leftIntegralSecond',
                                                                     type="text",
                                                                     debounce=True,
                                                                     min=-10000, max=10000, step=1,
                                                                     bs_size="sm",
                                                                     style={'width': '7rem', "marginTop": "1.5rem"},
                                                                     autoFocus=True,
                                                                     placeholder="second point"),
                                                           dbc.Input(id='leftIntegral',
                                                                     type="text",
                                                                     min=-10000, max=10000, step=1,
                                                                     bs_size="sm",
                                                                     style={'width': '8rem', "marginTop": "1.5rem"},
                                                                     autoFocus=True,
                                                                     placeholder="total integration"),
                                                           ]),
                                                 html.Div([html.Button("Save", id="write_excel", n_clicks=0,
                                                                       style={'fontSize': '1rem', 'width': '4rem',
                                                                              'margin': '1rem'},
                                                                       ),
                                                           html.A(html.Button("Download Data", id='download_data',
                                                                              n_clicks=0,
                                                                              style={'fontSize': '1rem',
                                                                                     'width': '9rem',
                                                                                     'margin': '1rem'}, ),
                                                                  id='download_excel',
                                                                  # # download="rawdata.csv",
                                                                  href="/download_excel/",
                                                                  # target="_blank"
                                                                  )
                                                           ], className='ad')

                                                 ]),
                                       html.Div([dbc.Checklist(
                                           id='operateur',
                                           options=[{'label': i, 'value': i} for i in
                                                    ['Plus', 'Moins', 'Multiplie', 'Division']],
                                           value=[],
                                           labelStyle={"display": "Block"},
                                       ), ]),
                                       html.Div([dcc.Dropdown(id='secondChoosenValue',
                                                              options=[{'label': i, 'value': i} for i in
                                                                       dropdownhidden],
                                                              multi=False,
                                                              style={"cursor": "pointer", 'width': '180px'},
                                                              className='',
                                                              clearable=True,
                                                              placeholder='Second Value...',
                                                              ),
                                                 dbc.Input(id='rightIntegralFirst',
                                                           type="text",
                                                           min=-10000, max=10000, step=1,
                                                           bs_size="sm",
                                                           style={'width': '7rem', "marginTop": "1.5rem"},
                                                           autoFocus=True,
                                                           placeholder="first point"),
                                                 dbc.Input(id='rightIntegralSecond',
                                                           type="text",
                                                           min=-10000, max=10000, step=1,
                                                           bs_size="sm",
                                                           style={'width': '7rem', "marginTop": "1.5rem"},
                                                           autoFocus=True,
                                                           placeholder="second point"),
                                                 dbc.Input(id='rightIntegral',
                                                           type="text",
                                                           min=-10000, max=10000, step=1,
                                                           bs_size="sm",
                                                           style={'width': '8rem', "marginTop": "1.5rem"},
                                                           autoFocus=True,
                                                           placeholder="total integration")]),
                                       html.Div([dbc.Input(id='operation',
                                                           type="text",
                                                           min=-10000, max=10000, step=1,
                                                           bs_size="sm",
                                                           style={'width': '10rem', "marginTop": "2rem",
                                                                  'height': '5rem', 'textAlign': 'center'},
                                                           autoFocus=True,
                                                           placeholder="result"),
                                                 dbc.Input(id='intersection',
                                                           type="text",
                                                           min=-10000, max=10000, step=1,
                                                           bs_size="sm",
                                                           style={'width': '10rem', "marginTop": "2rem",
                                                                  'height': '2rem', 'textAlign': 'center'},
                                                           autoFocus=True,
                                                           placeholder="Intersection")], className='aa')],
                                      className="ab"),
                             html.Div([dcc.RadioItems(id="radiograph",
                                                      options=[
                                                          {'label': 'Point', 'value': 'markers'},
                                                          {'label': 'Line', 'value': 'lines'},
                                                          {'label': 'Line + Point', 'value': 'lines+markers'},

                                                      ],
                                                      value='markers',
                                                      labelClassName='groupgraph2',
                                                      labelStyle={'margin': '10px', },
                                                      inputStyle={'margin': '10px', }
                                                      ),
                                       dbc.Input(id='minimumValueGraphFirst',
                                                 type="text",
                                                 min=-10000, max=10000, step=1,
                                                 bs_size="sm",
                                                 value=0,
                                                 style={'width': '7rem', "marginTop": "1rem"},
                                                 placeholder="Minimum Value of Graph for First..."),
                                       dbc.Input(id='minimumValueGraphSecond',
                                                 type="text",
                                                 min=-10000, max=10000, step=1,
                                                 bs_size="sm",
                                                 value=0,
                                                 style={'width': '7rem', "marginTop": "1rem", 'marginLeft': '0.5rem'},
                                                 placeholder="Minimum Value of Graph for Second..."),

                                       ], className='abcd'),

                             html.Div([dcc.Dropdown(id='shiftaxisdrop',
                                                    options=[{'label': i, 'value': i} for i in
                                                             dropdownhidden],
                                                    multi=False,
                                                    style={"cursor": "pointer", 'width': '180px', 'margin': '1rem'},
                                                    className='',
                                                    clearable=True,
                                                    placeholder='Choose Value...',
                                                    ),
                                       dbc.Button("See Surface", id="valuechange", n_clicks=0,
                                                  color="warning", style={'height': '2.5em', 'margin': '1.8rem'}),
                                       dbc.Button("Clean Surface", id="cleanshape", n_clicks=0,
                                                  color="danger", style={'height': '2.5em', 'margin': '1.8rem'}),
                                       html.Div(id='shiftaxis',
                                                children=[
                                                    dbc.Input(id='shift_x_axis',
                                                              type="number",
                                                              min=-100000, max=100000, step=1,
                                                              bs_size="sm",
                                                              value=0,
                                                              style={'width': '7rem', },
                                                              placeholder="Shift X axis..."),
                                                    dbc.Input(id='shift_y_axis',
                                                              type="number",
                                                              min=-100000, max=100000, step=1,
                                                              bs_size="sm",
                                                              value=0,
                                                              style={'width': '7rem', },
                                                              placeholder="Shift Y axis..."),
                                                    dbc.Button("Ok", id="tab1send", outline=True, n_clicks=0,
                                                               color="primary",
                                                               className="mr-2"),
                                                ], className='abcd', style={'display': 'None'}),

                                       ], className='abcd'),

                             html.Div([dcc.Graph(id='graph',
                                                 config={'displayModeBar': True,
                                                         'scrollZoom': True,
                                                         'modeBarButtonsToAdd': [
                                                             'drawline',
                                                             'drawrect',
                                                             'drawopenpath',
                                                             'select2d',
                                                             'eraseshape',
                                                         ]},
                                                 style={'marginTop': 20},
                                                 figure={
                                                     'layout': {'legend': {'tracegroupgap': 0},

                                                                }
                                                 }

                                                 ),
                                       html.Div(daq.Slider(id="sliderHeightTab1",
                                                           max=2100,
                                                           min=400,
                                                           value=530,
                                                           step=100,
                                                           size=400,
                                                           vertical=True,
                                                           updatemode='drag'), style={'margin': '20px'})],
                                      className='abcdb'),

                             html.Div([daq.Slider(id="sliderWidthTab1",
                                                  max=2000,
                                                  min=600,
                                                  value=1000,
                                                  step=100,
                                                  size=750,
                                                  updatemode='drag'),
                                       html.Div(id='output-data-upload', children=[])]),

                             ])

        return loadTab1


# bunu bi duzeltmeye calisacam
@app.callback(Output("leftSideChecklistValueHidden", "children"),
              [Input('choosenChecklistLeft', 'value'), ],
              [State("leftSideChecklistValueHidden", "children")]
              )
def res(val, hiddenval):
    if val == None:
        raise PreventUpdate
    print('val', val)
    hiddenval = val
    return hiddenval


# @app.callback(Output("leftSideChecklistValueHiddenTab4", "children"),
#               [Input('choosenChecklistLeft', 'value')],
#               [State("leftSideChecklistValueHiddenTab4", "children")]
#               )
# def res(val, hiddenval):
#     if val == None:
#         raise PreventUpdate
#     hiddenval = val
#     print('valllllllllll', val)
#     print('hiddenval', hiddenval)
#     return hiddenval


@app.callback(Output("radiographhidden", "children"),
              [Input("radiograph", "value")],
              )
def radio(radiograph):
    return radiograph


@app.callback(Output("radiographhiddentab4", "children"),
              [Input("radiograph4", "value")],
              )
def radiotab4(radiograph):
    if radiograph == None:
        raise PreventUpdate
    return radiograph


@app.callback(Output("sliderHeightTab1hidden", "children"),
              [Input("sliderHeightTab1", "value")],
              )
def tabheight(height):
    return height


@app.callback(Output("sliderWidthTab1hidden", "children"),
              [Input("sliderWidthTab1", "value")],
              )
def tabwidth(width):
    return width


@app.callback(Output("sliderHeightTab1hiddentab4", "children"),
              [Input("sliderHeightTab4", "value")],
              )
def tabheighttab4(height):
    return height


@app.callback(Output("sliderWidthTab1hiddenTab4", "children"),
              [Input("sliderWidthTab4", "value")],
              )
def tabwidthtab4(width):
    return width


@app.callback(Output("minimumValueGraphhiddenfirst", "children"),
              [Input("minimumValueGraphFirst", "value")],
              )
def minfirst(minvalfirst):
    return minvalfirst


@app.callback(Output("minimumValueGraphhiddensecond", "children"),
              [Input("minimumValueGraphSecond", "value")],
              )
def minsecond(minvalsecond):
    return minvalsecond


@app.callback(Output("firstchoosenvalhidden", "children"),
              [Input("firstChoosenValue", "value")],
              [State("firstchoosenvalhidden", "children")]
              )
def firstchleft(firstchoosen, hiddenfirstchoosen):
    hiddenfirstchoosen.append(firstchoosen)
    return hiddenfirstchoosen


@app.callback(Output("firstchoosenvalhiddentab4", "children"),
              [Input("firstChoosenValueTab4", "value")],
              [State("firstchoosenvalhiddentab4", "children")]
              )
def firstchlefttab4(firstchoosen4, hiddenfirstchoosen4):
    hiddenfirstchoosen4.append(firstchoosen4)
    return hiddenfirstchoosen4


@app.callback(Output("secondchoosenvalhidden", "children"),
              [Input("secondChoosenValue", "value")],
              )
def secondchleft(secondchoosen):
    return secondchoosen


@app.callback(Output("secondchoosenvalhiddentab4", "children"),
              [Input("secondChoosenValue", "value")],
              )
def secondchleft(secondchoosen):
    return secondchoosen


@app.callback(Output("leftintegralfirsthidden", "children"),
              [Input("leftIntegralFirst", "value")],
              )
def firstchright(leftintfirst):
    return leftintfirst


@app.callback(Output("leftintegralfirsthiddenTab4", "children"),
              [Input("leftIntegralFirstTab4", "value")],
              )
def firstchrighttab4(leftintfirst):
    return leftintfirst


@app.callback(Output("leftintegralsecondhidden", "children"),
              [Input("leftIntegralSecond", "value")],
              )
def secondchright(leftintsecond):
    return leftintsecond


@app.callback(Output("leftintegralsecondhiddentab4", "children"),
              [Input("leftIntegralSecondTab4", "value")],
              )
def secondchright(leftintsecond):
    return leftintsecond


@app.callback(Output("rightintegralfirsthidden", "children"),
              [Input("rightIntegralFirst", "value")],
              )
def rightfrst(rightintfirst):
    return rightintfirst


@app.callback(Output("rightintegralfirsthiddentab4", "children"),
              [Input("rightIntegralFirstTab4", "value")],
              )
def rightfrsttab4(rightintfirst):
    return rightintfirst


@app.callback(Output("rightintegralsecondhidden", "children"),
              [Input("rightIntegralSecond", "value")],
              )
def rightscnd(rightintsecond):
    return rightintsecond


@app.callback(Output("rightintegralsecondhiddentab4", "children"),
              [Input("rightIntegralSecondTab4", "value")],
              )
def rightscndtab4(rightintsecond):
    return rightintsecond


##### bunla ugras shapeler ciktiktan sonra referance bilgileri cikmiyor
@app.callback([Output("inputRightY_axishidden", "children"), Output("inputRightX_axishidden", "children"),
               ],
              [Input('valueSendRight', 'n_clicks')],
              [State("inputRightY_axis", "value"), State("inputRightX_axis", "value"),
               State("inputRightY_axishidden", "children"), State("inputRightX_axishidden", "children"),
               ]
              )
def Inputaxis(nclick, y_val, x_val, y_axis, x_axis):
    if nclick > 0:
        y_axis.append(y_val)
        x_axis.append(x_val)
        return (y_axis, x_axis)
    else:
        return no_update


# @app.callback(Output('valueSendRighthidden','children'),
#               [Input('valueSendRight','n_clicks')])
# def sendright(val):
#     return val
#
# @app.callback(Output('checklistvaleurhidden', "children"),
#               [Input('checklistValeur','value')])
# def sendrightdrop(val):
#     return val
# for show graph and changement

@app.callback(Output('shiftaxisdrophidden', 'children'),
              [Input('shiftaxisdrop', 'value')], )
def relay(val):
    return val


@app.callback(Output('shift_x_axishidden', 'children'),
              [Input('shift_x_axis', 'value')], )
def relay2(val):
    return val


@app.callback(Output('shift_y_axishidden', 'children'),
              [Input('shift_y_axis', 'value')], )
def relay3(val):
    return val


@app.callback(Output('tab1sendhidden', 'children'),
              [Input('tab1send', 'n_clicks')], )
def relay7(val):
    return val


@app.callback(Output('shiftaxis', 'style'),
              [Input('shiftaxisdrop', 'value')])
def shiftingaxes(val):
    if val == None:
        return {'display': 'None'}
    return {'visibility': 'visible', 'marginTop': '2rem'}


@app.callback([Output('graphhidden', 'figure'),
               Output('hiddenDifferance', 'children'), ],
              [Input("choosenChecklistLeft", "value"),
               Input("radiographhidden", "children"),
               Input("sliderHeightTab1hidden", "children"),
               Input("sliderWidthTab1hidden", "children"),
               Input('minimumValueGraphhiddenfirst', 'children'),
               Input('minimumValueGraphhiddensecond', 'children'),
               Input('firstchoosenvalhidden', 'children'),
               Input('secondchoosenvalhidden', 'children'),
               Input('checklistvaleurhidden', "children"),
               Input('inputRightY_axishidden', 'children'),
               Input('inputRightX_axishidden', 'children'),
               Input('tab1sendhidden', 'children'),
               Input('valuechange', 'n_clicks'),
               Input('cleanshape', 'n_clicks'),
               ],
              [State('shiftaxisdrophidden', 'children'),
               State('shift_x_axishidden', 'children'),
               State('shift_y_axishidden', 'children'),
               State('hiddenDifferance', 'children'),
               State('retrieve', 'children'),
               State('leftintegralfirsthidden', 'children'),
               State('leftintegralsecondhidden', 'children'),
               State('rightintegralfirsthidden', 'children'),
               State('rightintegralsecondhidden', 'children'),
               State('pointLeftFirst', 'children'),
               State('pointRightFirst', 'children'),
               ]
              )
def res2(val, radiograph, sliderheight, sliderwidth,
         minValfirst, minValsecond, firstchoosen, secondchoosen, rightsidedrop, right_y_axis, right_x_axis,
         nclick, nc, cleanclick, axis, shift_x, shift_y, differance, retrieve, leftfirstval, leftsecondval,
         rightfirstval, rightsecondval, firstshape, secondshape, ):
    if retrieve == None or retrieve == [] or val == []:
        raise PreventUpdate
    if len(retrieve) > 0:
        print('grapval', val)
        df = pd.read_excel('appending.xlsx')
        df['index'] = df.index
        df = df.reindex(columns=sorted(df.columns, reverse=True))
        baseval = ''
        if 'date' not in df.columns:
            for col in df.columns:
                if 'Temps' in col:
                    baseval += col
                    dt = df[baseval]
        else:
            df_shape = df.copy()
            df_shape['newindex'] = df_shape.index
            df_shape.index = df_shape['date']
            dt = ["{}-{:02.0f}-{:02.0f}_{:02.0f}:{:02.0f}:{:02.0f}".format(d.year, d.month, d.day, d.hour, d.minute,
                                                                           d.second) for d in df_shape.index]
        fig = go.Figure()
        if right_x_axis != [] and right_y_axis != []:
            for k in range(len(rightsidedrop)):
                if right_x_axis[k] != None or right_y_axis[k] != None:
                    x = int(right_x_axis[k])
                    y = int(right_y_axis[k])
                    z = int(right_x_axis[k]) / 100
                    t = int(right_y_axis[k]) / 100
                    fig.add_shape(type="circle",
                                  x0=x, y0=y, x1=x + z, y1=y + t,
                                  xref="x", yref="y",
                                  fillcolor="PaleTurquoise",
                                  line_color="LightSeaGreen",
                                  )
                    fig.add_annotation(x=x, y=y,
                                       text="{} - {} référence".format(x, y),
                                       showarrow=True,
                                       yshift=80
                                       )
                else:
                    no_update

        for i_val in range(len(val)):
            y_axis = df[val[i_val]]
            if 'date' not in df.columns:
                x_axis = df[baseval]
            else:
                x_axis = df['date']
            if nclick > 0:
                if axis == val[i_val]:
                    j = []
                    for i in df[axis]:
                        if shift_y == None:
                            raise PreventUpdate
                        else:
                            i += float(shift_y)
                            j.append(i)
                    df[axis] = pd.DataFrame(j)
                    y_axis = df[axis]
                    df.to_excel("appending.xlsx")

                    if 'date' not in df.columns:
                        p = []
                        for i in df[baseval]:
                            if shift_x == None:
                                raise PreventUpdate
                            else:
                                i += float(shift_x)
                                p.append(i)
                        df['New x-axis'] = pd.DataFrame(p)
                        x_axis = df['New x-axis']
                        df.to_excel("appending.xlsx")
                    else:
                        x_axis = df['date']

            fig.add_trace(go.Scattergl(x=x_axis, y=y_axis, mode=radiograph, name=val[i_val]))
            color = {0: 'blue', 1: 'red', 2: 'green', 3: 'purple', 4: 'orange'}
            if len(firstshape) == 2 and leftfirstval != firstshape[0] and leftfirstval != []:
                if leftfirstval.startswith('T') == 1:
                    del firstshape[0]
                    firstshape.append(float(leftfirstval[2:]))
                    firstshape = sorted(firstshape)
                elif leftfirstval.isnumeric() == 1:
                    del firstshape[0]
                    firstshape.append(float(leftfirstval))
                    firstshape = sorted(firstshape)
                elif leftfirstval != None:
                    del firstshape[0]
            if len(firstshape) == 2 and leftsecondval != firstshape[
                1] and leftsecondval != None and leftsecondval != []:
                if leftsecondval.startswith('T') == 1:
                    del firstshape[1]
                    firstshape.append(float(leftsecondval[2:]))
                    firstshape = sorted(firstshape)
                elif leftsecondval.isnumeric() == 1:
                    del firstshape[1]
                    firstshape.append(float(leftsecondval))
                    firstshape = sorted(firstshape)
                elif leftsecondval != None:
                    del firstshape[1]

            if len(secondshape) == 2 and rightfirstval != secondshape[
                0] and rightfirstval != None and rightfirstval != []:
                if rightfirstval.startswith('T') == 1:
                    del secondshape[0]
                    secondshape.append(float(rightfirstval[2:]))
                    secondshape = sorted(secondshape)
                elif rightfirstval.isnumeric() == 1:
                    del secondshape[0]
                    secondshape.append(float(rightfirstval))
                    secondshape = sorted(secondshape)
                elif rightfirstval != None:
                    del secondshape[0]
            if len(secondshape) == 2 and rightsecondval != secondshape[
                1] and rightsecondval != None and rightsecondval != []:
                if rightsecondval.startswith('T') == 1:
                    del secondshape[1]
                    secondshape.append(float(rightsecondval[2:]))
                    secondshape = sorted(secondshape)
                elif rightsecondval.isnumeric() == 1:
                    del secondshape[1]
                    secondshape.append(float(rightsecondval))
                    secondshape = sorted(secondshape)
                elif rightsecondval != None:
                    del secondshape[1]
            if len(secondshape) == 2 and secondchoosen == None:
                del secondshape[1]
            if len(firstshape) == 2 and firstchoosen == None:
                del firstshape[1]

            def controlShape():
                pathline = ''
                pathline2 = ''
                if firstchoosen[-1] != None and secondchoosen != None:
                    if len(firstshape) == 2 and leftfirstval != None and leftsecondval != None:
                        if int(firstshape[1]) > int(firstshape[0]):
                            pathline = ''
                            rangeshape = range(int(firstshape[0]), int(firstshape[1]))
                            if ':' or '-' in dt[0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        print('111111111')
                                        pathline += 'M ' + str(dt[k]) + ', ' + str(minValfirst) + ' L' + str(
                                            dt[k]) + ', ' + str(df[firstchoosen[-1]][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(dt[k]) + ', ' + str(df[firstchoosen[-1]][k])
                                pathline += ' L' + str(dt[k]) + ', ' + str(minValfirst)
                                pathline += ' Z'
                                print('2222222')
                            else:
                                print('333333')
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(int(dt[k])) + ', ' + str(minValfirst) + ' L' + str(
                                            int(dt[k])) + ', ' + str(
                                            df[firstchoosen[-1]][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(int(dt[k])) + ', ' + str(df[firstchoosen[-1]][k])
                                pathline += ' L' + str(int(dt[k])) + ', ' + str(minValfirst)
                                pathline += ' Z'

                    if len(secondshape) == 2 and rightsecondval != None and rightfirstval != None:
                        if int(secondshape[1]) > int(secondshape[0]):
                            rangeshape = range(int(secondshape[0]), int(secondshape[1]))
                            if ':' or '-' in dt[0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(dt[k]) + ', ' + str(minValsecond) + ' L' + str(
                                            dt[k]) + ', ' + str(
                                            df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(dt[k]) + ', ' + str(df[secondchoosen][k])
                                pathline2 += ' L' + str(dt[k]) + ', ' + str(minValsecond)
                                pathline2 += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(int(dt[k])) + ', ' + str(minValsecond) + ' L' + str(
                                            int(dt[k])) + ', ' + str(
                                            df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(int(dt[k])) + ', ' + str(df[secondchoosen][k])
                                pathline2 += ' L' + str(int(dt[k])) + ', ' + str(minValsecond)
                                pathline2 += ' Z'

                    return [dict(
                        type="path",
                        path=pathline,
                        layer='below',
                        fillcolor="#5083C7",
                        opacity=0.8,
                        line_color="#8896BF",
                    ), dict(
                        type="path",
                        path=pathline2,
                        layer='below',
                        fillcolor="#B0384A",
                        opacity=0.8,
                        line_color="#B36873",
                    )]

                if firstchoosen[-1] != None and secondchoosen == None:
                    if len(firstshape) == 2:
                        if int(firstshape[1]) > int(firstshape[0]):
                            pathline = ''
                            rangeshape = range(int(firstshape[0]), int(firstshape[1]))
                            print('rangeshape', rangeshape)
                            if ':' or '-' in dt[0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        print('vsdfsfssdfsdfdfs2')
                                        pathline += 'M ' + str(dt[k]) + ', ' + str(minValfirst) + ' L' + str(
                                            dt[k]) + ', ' + str(df[firstchoosen[-1]][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(dt[k]) + ', ' + str(df[firstchoosen[-1]][k])
                                pathline += ' L' + str(dt[k]) + ', ' + str(minValfirst)
                                pathline += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        print('vsdfsfssdfsdfdfs2')
                                        pathline += 'M ' + str(int(dt[k])) + ', ' + str(minValfirst) + ' L' + str(
                                            int(dt[k])) + ', ' + str(
                                            df[firstchoosen[-1]][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(int(dt[k])) + ', ' + str(df[firstchoosen[-1]][k])
                                pathline += ' L' + str(int(dt[k])) + ', ' + str(minValfirst)
                                pathline += ' Z'

                            return [dict(
                                type="path",
                                path=pathline,
                                layer='below',
                                fillcolor="#5083C7",
                                opacity=0.8,
                                line_color="#8896BF",
                            )]

                        if int(firstshape[0]) > int(firstshape[1]):
                            rangeshape = range(int(firstshape[1]), int(firstshape[0]))
                            if ':' or '-' in dt[0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(dt[k]) + ', ' + str(minValsecond) + ' L' + str(
                                            dt[k]) + ', ' + str(
                                            df[firstchoosen[-1]][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(dt[k]) + ', ' + str(df[firstchoosen[-1]][k])
                                pathline += ' L' + str(dt[k]) + ', ' + str(minValsecond)
                                pathline += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(int(dt[k])) + ', ' + str(minValsecond) + ' L' + str(
                                            int(dt[k])) + ', ' + str(
                                            df[firstchoosen[-1]][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(int(dt[k])) + ', ' + str(df[firstchoosen[-1]][k])
                                pathline += ' L' + str(int(dt[k])) + ', ' + str(minValsecond)
                                pathline += ' Z'

                            return [dict(
                                type="path",
                                path=pathline,
                                layer='below',
                                fillcolor="#5083C7",
                                opacity=0.8,
                                line_color="#8896BF",
                            )]

                if secondchoosen != None and firstchoosen[-1] == None:
                    if len(secondshape) == 2 and rightsecondval != None and rightfirstval != None:
                        if int(secondshape[1]) > int(secondshape[0]):
                            rangeshape = range(int(secondshape[0]), int(secondshape[1]))
                            if ':' or '-' in dt[0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(dt[k]) + ', ' + str(minValsecond) + ' L' + str(
                                            dt[k]) + ', ' + str(
                                            df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(dt[k]) + ', ' + str(df[secondchoosen][k])
                                pathline2 += ' L' + str(dt[k]) + ', ' + str(minValsecond)
                                pathline2 += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(int(dt[k])) + ', ' + str(minValsecond) + ' L' + str(
                                            int(dt[k])) + ', ' + str(
                                            df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(int(dt[k])) + ', ' + str(df[secondchoosen][k])
                                pathline2 += ' L' + str(int(dt[k])) + ', ' + str(minValsecond)
                                pathline2 += ' Z'

                            return [dict(
                                type="path",
                                path=pathline2,
                                layer='below',
                                fillcolor="#5083C7",
                                opacity=0.8,
                                line_color="#8896BF",
                            )]

                        if int(secondshape[0]) > int(secondshape[1]):
                            rangeshape = range(int(secondshape[1]), int(secondshape[0]))
                            for k in rangeshape:
                                if k == rangeshape[0]:
                                    pathline2 += 'M ' + str(dt[k]) + ', ' + str(minValsecond) + ' L' + str(
                                        dt[k]) + ', ' + str(
                                        df[secondchoosen][k]) + ' '

                                elif k != rangeshape[0] and k != rangeshape[-1]:
                                    pathline2 += ' L' + str(dt[k]) + ', ' + str(df[secondchoosen][k])
                            pathline2 += ' L' + str(dt[k]) + ', ' + str(minValsecond)
                            pathline2 += ' Z'
                        else:
                            rangeshape = range(int(secondshape[1]), int(secondshape[0]))
                            for k in rangeshape:

                                if k == rangeshape[0]:
                                    pathline2 += 'M ' + str(int(dt[k])) + ', ' + str(minValsecond) + ' L' + str(
                                        int(dt[k])) + ', ' + str(
                                        df[secondchoosen][k]) + ' '

                                elif k != rangeshape[0] and k != rangeshape[-1]:
                                    pathline2 += ' L' + str(int(dt[k])) + ', ' + str(df[secondchoosen][k])
                            pathline2 += ' L' + str(int(dt[k])) + ', ' + str(minValsecond)
                            pathline2 += ' Z'

                        return [dict(
                            type="path",
                            path=pathline2,
                            layer='below',
                            fillcolor="#5083C7",
                            opacity=0.8,
                            line_color="#8896BF",
                        )]

            a = []
            if nc > 0:
                a = controlShape()
            fig.update_layout(
                autosize=False,
                width=sliderwidth,
                height=sliderheight,
                shapes=a if nc > cleanclick else [],
                margin=dict(
                    l=50,
                    r=50,
                    b=100,
                    t=50,
                    pad=4
                ),
                paper_bgcolor="LightSteelBlue",
            )

            if len(firstshape) == 2 and len(secondshape) == 2:
                a = int(firstshape[0])
                c = int(secondshape[0])
                b = int(firstshape[1])
                d = int(secondshape[1])
                if len(set(range(a, b)).intersection(set(range(c, d)))) >= 1 or len(
                        set(range(c, d)).intersection(set(range(a, b)))) >= 1:
                    if a <= c:
                        if len(differance) == 2:
                            differance.pop(0)
                            differance.append(b)
                        differance.append(b)
                    if a >= c:
                        if len(differance) == 2:
                            differance.pop(0)
                            differance.append(a)
                        differance.append(a)
                    if b <= d:
                        if len(differance) == 2:
                            differance.pop(0)
                            differance.append(c)
                        differance.append(c)
                    if b >= d:
                        if len(differance) == 2:
                            differance.pop(0)
                            differance.append(d)
                        differance.append(d)
                    if set(range(a, b)).issuperset(set(range(c, d))) == 1:
                        differance.append(c)
                        differance.append(d)
                    if set(range(c, d)).issuperset(set(range(a, b))) == 1:
                        differance.append(a)
                        differance.append(b)

                else:
                    differance = [0, 0]
        return fig, differance[-2:]

    else:
        return (no_update, no_update)


@app.callback(Output('graph', 'figure'),
              [Input("graphhidden", "figure")], )
def aa(a):
    return a


# @app.callback(Output('tab2Data', 'children'),
#               [Input("my-toggle-switch", "on")],
#               )
# def LoadingDataTab2(on):
#
#     if on == 1:
#
#         data_list = ['CoAd', 'ComManCoP2', 'ComManCoP3P4P5', 'ComManPompeSec', 'CompteurEnergie', 'CoP2',
#                          'CtempDepChauff',
#                          'D1', 'D2', 'D3', 'D4', 'MarcheBruleur', 'Teg', 'SdeBasBouMelange', 'SdeBasHauMelange', 'TambN3',
#                          'Tb1',
#                          'Tb2', 'Tb3', 'Tb4', 'TdepPLC', 'Teb', 'Tec', 'Teev', 'TempminMaf', 'Text', 'Tsb', 'Tsc', 'Tsev']
#
#         loadlist = html.Div(children=[
#                 html.Div([html.Div([html.Div([dcc.Dropdown(id='tabDropdownTop',
#                                                            options=[{'label': i, 'value': i} for i in data_list],
#                                                            multi=True,
#                                                            style={"cursor": "pointer"},
#                                                            className='stockSelectorClass2',
#                                                            clearable=True,
#                                                            placeholder='Select your y-axis value...',
#                                                            ),
#                                               dcc.Dropdown(id='tabDropdownDown',
#                                                            options=[{'label': i, 'value': i} for i in data_list],
#                                                            multi=True,
#                                                            style={"cursor": "pointer"},
#                                                            className='stockSelectorClass2',
#                                                            clearable=True,
#                                                            placeholder='Select your x-axis value...',
#                                                            ), ], className="ab"),
#                                     html.Div(dcc.RadioItems(id="radiograph2",
#                                                             options=[
#                                                                 {'label': 'Point', 'value': 'markers'},
#                                                                 {'label': 'Line', 'value': 'lines'},
#                                                                 {'label': 'Line + Point', 'value': 'lines+markers'}],
#                                                             value='markers',
#                                                             labelClassName='groupgraph2',
#                                                             labelStyle={'margin': '10px', },
#                                                             inputStyle={'margin': '10px', }
#                                                             ), ), ], className="ac"),
#                           html.Div([dcc.Dropdown(id="dropadd",
#                                                  options=[
#                                                      {'label': 'Note', 'value': 'note'},
#                                                      {'label': 'Header', 'value': 'header'},
#                                                      {'label': 'x-axis', 'value': 'x_axis'},
#                                                      {'label': 'y-axis', 'value': 'y_axis'},
#
#                                                  ],
#                                                  value='header',
#                                                  ),
#                                     dcc.Textarea(
#                                         id='textarea',
#                                         value='',
#                                         style={'width': '15rem', 'marginTop': '0.5rem'},
#                                         autoFocus='Saisir',
#                                     ),
#                                     ], className="aa"),
#
#                           html.Button('addText', id='addText', n_clicks=0, style={'marginTop': '1.5rem'}),
#
#                           ], className="tabDesign", ),
#
#                 html.Div([dcc.Graph(id='graph2', config={'displayModeBar': True,
#                                                          'scrollZoom': True,
#                                                          'modeBarButtonsToAdd': [
#                                                              'drawopenpath',
#                                                              'drawcircle',
#                                                              'eraseshape',
#                                                              'select2d',
#                                                          ]},
#                                     figure={
#                                         'layout': {'legend': {'tracegroupgap': 0},
#
#                                                    }
#                                     }
#                                     ),
#                           dcc.Slider(id="sliderHeight",
#                                      max=2100,
#                                      min=400,
#                                      value=500,
#                                      step=100,
#                                      vertical=True,
#                                      updatemode='drag')], className='abc'),
#
#                 html.Div([dcc.Slider(id="sliderWidth",
#                                      max=2000,
#                                      min=600,
#                                      value=950,
#                                      step=100,
#                                      updatemode='drag'),
#                           html.Div(id="tab2DashTable", children=[])]),
#             ])
#
#         return loadlist

@app.callback(Output('tab4Data', 'children'),
              [Input("my-toggle-switch", "on")],
              [State('tabs-with-classes', 'value')]
              )
def LoadingDataTab4(on, tab):
    if on == 1 and tab == 'tab-4':

        data_list = ['Choose your value firstly']

        loadlist = html.Div([html.Div([
            html.Div(id='tab4first', children=[html.Div([html.Div([
                dcc.RadioItems(id="radiographtab4",
                               options=[
                                   {'label': 'X - Y illimité', 'value': 'optionlibre'},
                                   {'label': 'Choose Values', 'value': 'choosevalue'},
                               ],
                               # value='choosevalue',
                               labelClassName='groupgraph',
                               labelStyle={'margin': '10px', },
                               inputStyle={'margin': '10px', }
                               ),

                dcc.RadioItems(id="radiograph4",
                               options=[
                                   {'label': 'Point', 'value': 'markers'},
                                   {'label': 'Line', 'value': 'lines'},
                                   {'label': 'Line + Point', 'value': 'lines+markers'}],
                               value='markers',
                               labelClassName='groupgraph2',
                               labelStyle={'margin': '10px', },
                               inputStyle={'margin': '10px', }
                               ),
            ], className="abtab4"),
                html.Div([dcc.Dropdown(id='tabDropdownTop4',
                                       options=[{'label': i, 'value': i} for i in data_list],
                                       multi=True,
                                       style={"cursor": "pointer", 'display': 'None'},
                                       className='stockSelectorClass2',
                                       clearable=True,
                                       placeholder='Select your y-axis value...',
                                       ),
                          dcc.Dropdown(id='tabDropdownTopTab4',
                                       options=[{'label': i, 'value': i} for i in data_list],
                                       multi=True,
                                       style={"cursor": "pointer", 'display': 'None'},
                                       className='stockSelectorClass2',
                                       clearable=True,
                                       placeholder='Select your y-axis value...',
                                       ),
                          dcc.Dropdown(id='tabDropdownDownTab4',
                                       options=[{'label': i, 'value': i} for i in data_list],
                                       multi=True,
                                       style={"cursor": "pointer", 'display': 'None'},
                                       className='stockSelectorClass2',
                                       clearable=True,
                                       placeholder='Select your x-axis value...',
                                       ),
                          dcc.Dropdown(id='tabDropdownTop',
                                       options=[{'label': i, 'value': i} for i in data_list],
                                       multi=True,
                                       style={"cursor": "pointer", 'display': 'None'},
                                       className='stockSelectorClass2',
                                       clearable=True,
                                       placeholder='Select your y-axis value...',
                                       ),
                          dcc.Dropdown(id='tabDropdownDown',
                                       options=[{'label': i, 'value': i} for i in data_list],
                                       multi=True,
                                       style={"cursor": "pointer", 'display': 'None'},
                                       className='stockSelectorClass2',
                                       clearable=True,
                                       placeholder='Select your x-axis value...',
                                       ),
                          ], className="ab"),
                html.Div([
                    dbc.Checklist(id="calculintegraltab4",
                                  options=[{'label': "Calculate Integral", 'value': 'calcultab4'}, ]
                                  ,
                                  value='',
                                  labelClassName='groupgraph',
                                  labelStyle={'margin': '10px', },
                                  inputStyle={'margin': '10px', }),
                ]), ], className="ac"),

                html.Div([dcc.Dropdown(id="dropadd4",
                                       options=[
                                           {'label': 'Note', 'value': 'note'},
                                           {'label': 'Header', 'value': 'header'},
                                           {'label': 'x-axis', 'value': 'x_axis'},
                                           {'label': 'y-axis', 'value': 'y_axis'},

                                       ],
                                       value='header',
                                       ),
                          dcc.Textarea(
                              id='textarea4',
                              value='',
                              style={'width': '15rem', 'marginTop': '0.5rem'},
                              autoFocus='Saisir',
                          ),
                          ], className="aatab4"),

                html.Button('Add Text', id='addText4', n_clicks=0, style={'marginTop': '1.5rem', 'marginLeft': '2rem'}),

            ], className="tabDesigntab4", ),
            html.Div(id='tab4check', children=
            [html.Div([html.Div([dcc.Dropdown(id='firstChoosenValueTab4',
                                              options=[{'label': i, 'value': i} for i in
                                                       data_list],
                                              multi=False,
                                              style={"cursor": "pointer", 'width': '180px'},
                                              className='',
                                              clearable=True,
                                              placeholder='First Value...',
                                              ),
                                 dbc.Input(id='leftIntegralFirstTab4',
                                           type="text",
                                           debounce=True,
                                           min=-10000, max=10000, step=1,
                                           bs_size="sm",
                                           style={'width': '7rem', "marginTop": "1.5rem"},
                                           autoFocus=True,
                                           placeholder="first point"),
                                 dbc.Input(id='leftIntegralSecondTab4',
                                           type="text",
                                           debounce=True,
                                           min=-10000, max=10000, step=1,
                                           bs_size="sm",
                                           style={'width': '7rem', "marginTop": "1.5rem"},
                                           autoFocus=True,
                                           placeholder="second point"),
                                 dbc.Input(id='leftIntegralTab4',
                                           type="text",
                                           min=-10000, max=10000, step=1,
                                           bs_size="sm",
                                           style={'width': '8rem', "marginTop": "1.5rem"},
                                           autoFocus=True,
                                           placeholder="total integration"),
                                 ]), html.Div([html.Button("Save", id="write_excelTab4", n_clicks=0,
                                                           style={'fontSize': '1rem', 'width': '4rem',
                                                                  'margin': '1rem'},
                                                           ),
                                               html.A(html.Button("Download Data", id='download_dataTab4',
                                                                  n_clicks=0,
                                                                  style={'fontSize': '1rem', 'width': '9rem',
                                                                         'margin': '1rem'}, ),
                                                      id='download_excelTab4',
                                                      # # download="rawdata.csv",
                                                      href="/download_excel/",
                                                      # target="_blank"
                                                      )
                                               ], className='ad')

                       ]),
             html.Div([dbc.Checklist(
                 id='operateurTab4',
                 options=[{'label': i, 'value': i} for i in
                          ['Plus', 'Moins', 'Multiplie', 'Division']],
                 value=[],
                 labelStyle={"display": "Block"},
             ), ]),
             html.Div([
                 dcc.Dropdown(id='secondChoosenValueTab4',
                              options=[{'label': i, 'value': i} for i in
                                       data_list],
                              multi=False,
                              style={"cursor": "pointer", 'width': '180px'},
                              className='',
                              clearable=True,
                              placeholder='Second Value...',
                              ),
                 dbc.Input(id='rightIntegralFirstTab4',
                           type="text",
                           min=-10000, max=10000, step=1,
                           bs_size="sm",
                           style={'width': '7rem', "marginTop": "1.5rem"},
                           autoFocus=True,
                           placeholder="first point"),
                 dbc.Input(id='rightIntegralSecondTab4',
                           type="text",
                           min=-10000, max=10000, step=1,
                           bs_size="sm",
                           style={'width': '7rem', "marginTop": "1.5rem"},
                           autoFocus=True,
                           placeholder="second point"),
                 dbc.Input(id='rightIntegralTab4',
                           type="text",
                           min=-10000, max=10000, step=1,
                           bs_size="sm",
                           style={'width': '8rem', "marginTop": "1.5rem"},
                           autoFocus=True,
                           placeholder="total integration")
             ]),
             html.Div([dbc.Input(id='operationTab4',
                                 type="text",
                                 min=-10000, max=10000, step=1,
                                 bs_size="sm",
                                 style={'width': '10rem', "marginTop": "2rem",
                                        'height': '5rem', 'textAlign': 'center'},
                                 autoFocus=True,
                                 placeholder="result"),
                       dbc.Input(id='intersectionTab4',
                                 type="text",
                                 min=-10000, max=10000, step=1,
                                 bs_size="sm",
                                 style={'width': '10rem', "marginTop": "2rem",
                                        'height': '2rem', 'textAlign': 'center'},
                                 autoFocus=True,
                                 placeholder="Intersection")], className='aa')
             ], style={'display': 'None'},
                     className="abTab4"),

            html.Div(id='tab4second', children=[dcc.Dropdown(id='shiftaxisdroptab4',
                                                             options=[{'label': i, 'value': i} for i in
                                                                      []],
                                                             multi=False,
                                                             style={"cursor": "pointer", 'width': '180px',
                                                                    'margin': '1rem'},
                                                             className='',
                                                             clearable=True,
                                                             placeholder='Choose Value...',
                                                             ),
                                                dbc.Button("See Surface", id="valuechangetab4", n_clicks=0,
                                                           color="warning",
                                                           style={'height': '2.5em', 'margin': '1.8rem'}),
                                                dbc.Button("Clean Surface", id="cleanshapetab4", n_clicks=0,
                                                           color="danger",
                                                           style={'height': '2.5em', 'margin': '1.8rem'}),
                                                html.Div(id='shiftaxistab4',
                                                         children=[
                                                             dbc.Input(id='shift_x_axistab4',
                                                                       type="number",
                                                                       min=-100000, max=100000, step=1,
                                                                       bs_size="sm",
                                                                       value=0,
                                                                       style={'width': '7rem', },
                                                                       placeholder="Shift X axis..."),
                                                             dbc.Input(id='shift_y_axistab4',
                                                                       type="number",
                                                                       min=-100000, max=100000, step=1,
                                                                       bs_size="sm",
                                                                       value=0,
                                                                       style={'width': '7rem', },
                                                                       placeholder="Shift Y axis..."),
                                                             dbc.Button("Ok", id="tab4send", outline=True, n_clicks=0,
                                                                        color="primary",
                                                                        className="mr-2"),
                                                         ], className='abcd',
                                                         style={'display': 'None'})

                                                ], className='abcd'),

            html.Div(id='tab4third', children=[dcc.Store(id='tab4datastore'),
                                               dcc.Graph(id='graph4', config={'displayModeBar': True,
                                                                              'scrollZoom': True,
                                                                              'modeBarButtonsToAdd': [
                                                                                  'drawopenpath',
                                                                                  'drawcircle',
                                                                                  'eraseshape',
                                                                                  'select2d',
                                                                              ]},
                                                         figure={
                                                             'layout': {'legend': {'tracegroupgap': 0},

                                                                        }
                                                         }
                                                         ),
                                               html.Div(daq.Slider(id="sliderHeightTab4",
                                                                   max=2100,
                                                                   min=400,
                                                                   value=800,
                                                                   step=100,
                                                                   size=400,
                                                                   vertical=True,
                                                                   updatemode='drag'), style={'margin': '10px'})],
                     className='abcTab4'),

            html.Div([daq.Slider(id="sliderWidthTab4",
                                 max=2000,
                                 min=600,
                                 value=1500,
                                 step=100,
                                 size=750,
                                 updatemode='drag'),
                      html.Div(id="tab4DashTable", children=[])
                      ]),
        ]), ])

        return loadlist
    else:
        no_update


@app.callback([Output('fourcolumnsdivusercontrols', 'style'),
               Output('eightcolumnsdivforcharts', 'style'), ],
              # Output('tab4third', 'style'),],
              Input('tabs-with-classes', 'value'), )
def tab4enlarger(tab):
    if tab == 'tab-4':
        return {'display': 'None'}, {'width': '260%', 'margin': '1rem'}
    if tab == 'tab-3':
        return {'display': 'None'}, {'width': '260%', 'margin': '1rem'}
    else:
        return {'visibility': 'visible'}, {'visibility': 'visible'}


@app.callback(Output('tab4check', 'style'),
              [Input("calculintegraltab4", "value")],
              )
def showintegral(show):
    if show == ['calcultab4']:
        return {'visibility': 'visible'}
    return {'display': 'None'}


@app.callback([Output("tabDropdownTop", "options"), Output("tabDropdownDown", "options")],
              [Input("retrieve", "children")])
def dropdownlistcontrol(retrieve):
    if len(retrieve) > 0:
        df = pd.read_excel('appending.xlsx')
        dff = [{'label': i, 'value': i} for i in df.columns if i.startswith('Un') != 1 and i != 'index' and i != 'date']
        return (dff, dff)
    else:
        return (no_update, no_update)


@app.callback([Output("tabDropdownTopTab4", "options"), Output("tabDropdownDownTab4", "options")],
              [Input("retrieve", "children")])
def dropdownlistcontrolTab4Second(retrieve):
    if len(retrieve) > 0:
        df = pd.read_excel('appending.xlsx')
        dff = [{'label': i, 'value': i} for i in df.columns if i.startswith('Un') != 1 and i != 'index' and i != 'date']
        return (dff, dff)
    else:
        return (no_update, no_update)


@app.callback([Output('tabDropdownTopTab4', 'style'),
               Output('tabDropdownDownTab4', 'style'),
               Output('tabDropdownTop', 'style'),
               Output('tabDropdownDown', 'style')
               ],
              [Input('radiographtab4', 'value')], )
def chooseradio(radio):
    if radio == None:
        raise PreventUpdate
    if radio == 'choosevalue':
        return {'visibility': 'visible'}, {'visibility': 'visible'}, {'display': 'None'}, {'display': 'None'}
    if radio == 'optionlibre':
        return {'display': 'None'}, {'display': 'None'}, {'visibility': 'visible'}, {'visibility': 'visible'},


@app.callback([Output('tab2hiddenValuex_axis', 'children'),
               Output('tab2hiddenValuey_axis', 'children')],
              [Input('tabDropdownTop', 'value'),
               Input('tabDropdownDown', 'value'),
               Input('radiographtab4', 'value')],
              )
def contractdropdown(x, y, radioval):
    if x == [] or x == None or y == None or y == []:
        raise PreventUpdate
    if radioval == 'optionlibre':
        return x, y
    else:
        return [], []


@app.callback([Output('tab4hiddenValuex_axissecond', 'children'),
               Output('tab4hiddenValuey_axissecond', 'children'),
               ],
              [Input('tabDropdownTopTab4', 'value'),
               Input('tabDropdownDownTab4', 'value'),
               Input('radiographtab4', 'value')]
              )
def contractdropdown2(valxsecond, valysecond, radio):
    if valxsecond == None or valysecond == None or radio == None:
        raise PreventUpdate

    if radio == 'choosevalue':
        return valxsecond, valysecond

    else:
        return [], []


@app.callback(Output("tabDropdownTop4", "options"),
              [Input("retrieve", "children")])
def dropdownlistcontrolTab4First(retrieve):
    if len(retrieve) > 0:
        df = pd.read_excel('appending.xlsx')
        dff = [{'label': i, 'value': i} for i in df.columns if i.startswith('TG') == 1 or i[-2:] != 'ts']
        return dff
    else:
        return no_update


@app.callback(
    Output('output_s', 'children'),
    [Input('tabDropdownTopTab4', 'value'),
     Input('tabDropdownTop', 'value'),
     Input('radiographtab4', 'value')], )
def container4(val2, val3, radio):
    if val2 == None and val3 == None or radio == None:
        raise PreventUpdate

    a = ''

    if radio == 'choosevalue':
        if val2 != None:
            a = val2
            return a
        else:
            return ''

    if radio == 'optionlibre':
        if val3 != None:
            a = val3
            return a
        else:
            return ''


@app.callback(
    Output('shiftaxisdroptab4', 'options'),
    [Input('tabDropdownTopTab4', 'value'),
     Input('tabDropdownTop', 'value'),
     Input('radiographtab4', 'value')], )
def container5(val2, val3, radio):
    if val2 == None and val3 == None or radio == None:
        raise PreventUpdate

    a = []

    if radio == 'choosevalue':
        if val2 != None:
            a = val2

    if radio == 'optionlibre':
        if val3 != None:
            a = val3
    return [{'label': i, 'value': i} for i in a]


# @app.callback(
#       [Output('firstChoosenValueTab4', 'value'),
#        Output('secondChoosenValueTab4', 'value'),],
#       [Input('radiographtab4', 'value')],)
#
# def clearbox(radioval) :
#     if radioval == 'choosevalue' or radioval == 'optionlibre' or radioval == 'Standart':
#         return '',''
#     else : raise PreventUpdate


@app.callback(
    [Output('firstChoosenValueTab4', 'options'),
     Output('secondChoosenValueTab4', 'options')],
    [Input('output_s', 'children'),
     Input('radiographtab4', 'value')], )
def container4_2(val, radio):
    if val == None or val == []:
        raise PreventUpdate
    a = []
    if radio == 'choosevalue':
        print('vallllllllll output olan2', val)
        a = [{'label': i, 'value': i} for i in val], [{'label': i, 'value': i} for i in val]
    elif radio == 'optionlibre':
        print('vallllllllll output olan3', val)
        a = [{'label': i, 'value': i} for i in val], [{'label': i, 'value': i} for i in val]
    print('son radioya gore optionslar', val)
    return a


@app.callback([Output('hiddenTextxaxis', 'children'), Output('hiddenTextyaxis', 'children'),
               Output('hiddenTextHeader', 'children'), Output('hiddenTextNote', 'children')],
              [Input('addText', 'n_clicks')],
              [State('textarea', 'value'), State('dropadd', 'value'),
               State('hiddenTextxaxis', 'children'), State('hiddenTextyaxis', 'children'),
               State('hiddenTextHeader', 'children'), State('hiddenTextNote', 'children')]
              )
def detailedGraph(addtextclick, textarea, add, g1, g2, head, note):
    if add == None or g1 == None or g2 == None or head == None or note == None:
        raise PreventUpdate

    if addtextclick > 0:
        if add == 'x_axis':
            g1.append(textarea)

        if add == 'y_axis':
            g2.append(textarea)

        if add == 'header':
            head.append(textarea)

        if add == 'note':
            note.append(textarea)
        textarea = ''
        return g1, g2, head, note
    else:
        return (no_update, no_update, no_update, no_update)


@app.callback([Output('hiddenTextxaxis4', 'children'), Output('hiddenTextyaxis4', 'children'),
               Output('hiddenTextHeader4', 'children'), Output('hiddenTextNote4', 'children')],
              [Input('addText4', 'n_clicks')],
              [State('textarea4', 'value'), State('dropadd4', 'value'),
               State('hiddenTextxaxis4', 'children'), State('hiddenTextyaxis4', 'children'),
               State('hiddenTextHeader4', 'children'), State('hiddenTextNote4', 'children')]
              )
def detailedGraph4(addtextclick, textarea, add, g1, g2, head, note):
    if add == None or g1 == None or g2 == None or head == None or note == None:
        raise PreventUpdate

    if addtextclick > 0:
        if add == 'x_axis':
            g1.append(textarea)

        if add == 'y_axis':
            g2.append(textarea)

        if add == 'header':
            head.append(textarea)

        if add == 'note':
            note.append(textarea)
        textarea = ''
        return g1, g2, head, note
    else:
        return (no_update, no_update, no_update, no_update)


@app.callback(Output('shiftaxistab4', 'style'),
              [Input('shiftaxisdroptab4', 'value')])
def shiftingaxestab4(val):
    if val == None:
        return {'display': 'None'}
    return {'visibility': 'visible', 'marginTop': '2rem'}


@app.callback(Output('shiftaxisdroptab4hidden', 'children'),
              [Input('shiftaxisdroptab4', 'value')], )
def relay4(val):
    return val


@app.callback(Output('shift_x_axistab4hidden', 'children'),
              [Input('shift_x_axistab4', 'value')], )
def relay5(val):
    return val


@app.callback(Output('shift_y_axistab4hidden', 'children'),
              [Input('shift_y_axistab4', 'value')], )
def relay6(val):
    return val


@app.callback(Output('radiographtab4hidden', 'children'),
              [Input('radiographtab4', 'value')], )
def relay7(valradio):
    if valradio == None:
        raise PreventUpdate
    return valradio


@app.callback(Output('graph4', 'figure'),
              [Input('radiograph4', 'value'),
               Input('radiographtab4hidden', 'children'),
               Input('tab4hiddenValuex_axis', 'children'),
               Input('tab4hiddenValuex_axissecond', 'children'),
               Input('tab4hiddenValuey_axissecond', 'children'),
               Input('sliderHeightTab4', 'value'),
               Input('sliderWidthTab4', 'value'),
               Input('hiddenTextxaxis4', 'children'),
               Input('hiddenTextyaxis4', 'children'),
               Input('hiddenTextHeader4', 'children'),
               Input('hiddenTextNote4', 'children'),
               Input('tab4send', 'n_clicks'),
               Input('firstChoosenValueTab4', 'value'),
               Input('secondChoosenValueTab4', 'value'),
               Input('valuechangetab4', 'n_clicks'),
               Input('tab2hiddenValuex_axis', 'children'),
               Input('tab2hiddenValuey_axis', 'children'),
               Input('cleanshapetab4', 'n_clicks'),
               ],
              [State('shiftaxisdroptab4hidden', 'children'),
               State('shift_x_axistab4hidden', 'children'),
               State('shift_y_axistab4hidden', 'children'),
               State('retrieve', 'children'),
               State('pointLeftFirstTab4', 'children'),
               State('pointRightFirstTab4', 'children'),
               State('leftIntegralFirstTab4', 'value'),
               State('leftIntegralSecondTab4', 'value'),
               State('rightIntegralFirstTab4', 'value'),
               State('rightIntegralSecondTab4', 'value'),
               ]
              )
def detailedGraph4(radio, radioval, valx, valxsecond, valysecond,
                   slideheight, slidewidth, g1, g2, head, note, nclick, firstchoosen, secondchoosen, nc,
                   valx2, valy2, cleanclick, axisdrop, shift_x, shift_y, retrieve, firstshape, secondshape,
                   leftfirstval, leftsecondval, rightfirstval, rightsecondval, ):
    if g1 == None or g2 == None or head == None or note == None or radioval == []:
        raise PreventUpdate
    print('firstchoosen', firstchoosen)
    if radioval != None:
        if len(retrieve) > 0:
            df = pd.read_excel("appending.xlsx")
            df.dropna(axis=0, inplace=True)
            fig = go.Figure()
            print('firstshape ne olmali', firstshape)

            def controlShape():
                pathline = ''
                pathline2 = ''
                minValfirst = 0
                minValsecond = 0
                if firstchoosen != None and secondchoosen != None:
                    if len(firstshape) == 2 and leftfirstval != None and leftsecondval != None:
                        if int(firstshape[1]) > int(firstshape[0]):
                            pathline = ''
                            rangeshape = range(int(firstshape[0]), int(firstshape[1]))
                            if ':' or '-' in df[lst[i][0]][0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(minValfirst) + ' L' + \
                                                    str(df[lst[i][0]][k]) + ', ' + str(df[firstchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(df[firstchoosen][k])
                                pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValfirst)
                                pathline += ' Z'
                            else:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(minValfirst) + ' L' + \
                                                    str(int(df[lst[i][0]][k])) + ', ' + str(df[firstchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(df[firstchoosen][k])
                                pathline += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(minValfirst)
                                pathline += ' Z'

                    if len(secondshape) == 2 and rightsecondval != None and rightfirstval != None:
                        if int(secondshape[1]) > int(secondshape[0]):
                            rangeshape = range(int(secondshape[0]), int(secondshape[1]))
                            if ':' or '-' in df[lst[i][0]][0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond) + ' L' + \
                                                     str(df[lst[i][0]][k]) + ', ' + str(df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(df[secondchoosen][k])
                                pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond)
                                pathline2 += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                            minValsecond) + ' L' + \
                                                     str(int(df[lst[i][0]][k])) + ', ' + str(df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(int(a[k])) + ', ' + str(df[secondchoosen][k])
                                pathline2 += ' L' + str(int(a[k])) + ', ' + str(minValsecond)
                                pathline2 += ' Z'

                    return [dict(
                        type="path",
                        path=pathline,
                        layer='below',
                        fillcolor="#5083C7",
                        opacity=0.8,
                        line_color="#8896BF",
                    ), dict(
                        type="path",
                        path=pathline2,
                        layer='below',
                        fillcolor="#B0384A",
                        opacity=0.8,
                        line_color="#B36873",
                    )]

                if firstchoosen != None and secondchoosen == None:
                    if len(firstshape) == 2:
                        if int(firstshape[1]) > int(firstshape[0]):
                            pathline = ''
                            rangeshape = range(int(firstshape[0]), int(firstshape[1]))
                            if ':' or '-' in df[lst[i][0]][0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(minValfirst) + ' L' + \
                                                    str(df[lst[i][0]][k]) + ', ' + str(df[firstchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(df[firstchoosen][k])
                                pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValfirst)
                                pathline += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                            minValfirst) + ' L' + str(
                                            int(df[lst[i][0]][k])) + ', ' + str(
                                            df[firstchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(int(a[k])) + ', ' + str(df[firstchoosen][k])
                                pathline += ' L' + str(int(a[k])) + ', ' + str(minValfirst)
                                pathline += ' Z'

                            return [dict(
                                type="path",
                                path=pathline,
                                layer='below',
                                fillcolor="#5083C7",
                                opacity=0.8,
                                line_color="#8896BF",
                            )]

                        if int(firstshape[0]) > int(firstshape[1]):
                            rangeshape = range(int(firstshape[1]), int(firstshape[0]))
                            if ':' or '-' in df[lst[i][0]][0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(
                                            minValsecond) + ' L' + str(
                                            df[lst[i][0]][k]) + ', ' + str(
                                            df[firstchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(
                                            df[firstchoosen][k])
                                pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond)
                                pathline += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        pathline += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                            minValsecond) + ' L' + \
                                                    str(int(df[lst[i][0]][k])) + ', ' + str(df[firstchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(df[firstchoosen][k])
                                pathline += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(minValsecond)
                                pathline += ' Z'

                            return [dict(
                                type="path",
                                path=pathline,
                                layer='below',
                                fillcolor="#5083C7",
                                opacity=0.8,
                                line_color="#8896BF",
                            )]

                if secondchoosen != None and firstchoosen == None:
                    if len(secondshape) == 2 and rightsecondval != None and rightfirstval != None:
                        if int(secondshape[1]) > int(secondshape[0]):
                            rangeshape = range(int(secondshape[0]), int(secondshape[1]))
                            if ':' or '-' in df[lst[i][0]][0]:
                                for k in rangeshape:
                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond) + ' L' + \
                                                     str(df[lst[i][0]][k]) + ', ' + str(df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(df[secondchoosen][k])
                                pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond)
                                pathline2 += ' Z'
                            else:
                                for k in rangeshape:

                                    if k == rangeshape[0]:
                                        pathline2 += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                            minValsecond) + ' L' + \
                                                     str(int(df[lst[i][0]][k])) + ', ' + str(df[secondchoosen][k]) + ' '

                                    elif k != rangeshape[0] and k != rangeshape[-1]:
                                        pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                            df[secondchoosen][k])
                                pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(minValsecond)
                                pathline2 += ' Z'

                            return [dict(
                                type="path",
                                path=pathline2,
                                layer='below',
                                fillcolor="#5083C7",
                                opacity=0.8,
                                line_color="#8896BF",
                            )]

                        if int(secondshape[0]) > int(secondshape[1]):
                            rangeshape = range(int(secondshape[1]), int(secondshape[0]))
                            for k in rangeshape:
                                if k == rangeshape[0]:
                                    pathline2 += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(
                                        minValsecond) + ' L' + str(
                                        df[lst[i][0]][k]) + ', ' + str(
                                        df[secondchoosen][k]) + ' '

                                elif k != rangeshape[0] and k != rangeshape[-1]:
                                    pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(df[secondchoosen][k])
                            pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond)
                            pathline2 += ' Z'
                        else:
                            rangeshape = range(int(secondshape[1]), int(secondshape[0]))
                            for k in rangeshape:

                                if k == rangeshape[0]:
                                    pathline2 += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(minValsecond) + ' L' + \
                                                 str(int(df[lst[i][0]][k])) + ', ' + str(df[secondchoosen][k]) + ' '

                                elif k != rangeshape[0] and k != rangeshape[-1]:
                                    pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(df[secondchoosen][k])
                            pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(minValsecond)
                            pathline2 += ' Z'

                        return [dict(
                            type="path",
                            path=pathline2,
                            layer='below',
                            fillcolor="#5083C7",
                            opacity=0.8,
                            line_color="#8896BF",
                        )]
                else:
                    return no_update

            if len(firstshape) == 2 and leftfirstval != firstshape[0] and leftfirstval != []:
                if leftfirstval.startswith('T') == 1:
                    del firstshape[0]
                    firstshape.append(float(leftfirstval[2:]))
                    firstshape = sorted(firstshape)
                elif leftfirstval.isnumeric() == 1:
                    del firstshape[0]
                    firstshape.append(float(leftfirstval))
                    firstshape = sorted(firstshape)
                elif leftfirstval != None:
                    del firstshape[0]
            if len(firstshape) == 2 and leftsecondval != firstshape[
                1] and leftsecondval != None and leftsecondval != []:
                if leftsecondval.startswith('T') == 1:
                    del firstshape[1]
                    firstshape.append(float(leftsecondval[2:]))
                    firstshape = sorted(firstshape)
                elif leftsecondval.isnumeric() == 1:
                    del firstshape[1]
                    firstshape.append(float(leftsecondval))
                    firstshape = sorted(firstshape)
                elif leftsecondval != None:
                    del firstshape[1]

            if len(secondshape) == 2 and rightfirstval != secondshape[
                0] and rightfirstval != None and rightfirstval != []:
                if rightfirstval.startswith('T') == 1:
                    del secondshape[0]
                    secondshape.append(float(rightfirstval[2:]))
                    secondshape = sorted(secondshape)
                elif rightfirstval.isnumeric() == 1:
                    del secondshape[0]
                    secondshape.append(float(rightfirstval))
                    secondshape = sorted(secondshape)
                elif rightfirstval != None:
                    del secondshape[0]
            if len(secondshape) == 2 and rightsecondval != secondshape[
                1] and rightsecondval != None and rightsecondval != []:
                if rightsecondval.startswith('T') == 1:
                    del secondshape[1]
                    secondshape.append(float(rightsecondval[2:]))
                    secondshape = sorted(secondshape)
                elif rightsecondval.isnumeric() == 1:
                    del secondshape[1]
                    secondshape.append(float(rightsecondval))
                    secondshape = sorted(secondshape)
                elif rightsecondval != None:
                    del secondshape[1]
            if len(secondshape) == 2 and secondchoosen == None:
                del secondshape[1]
            if len(firstshape) == 2 and firstchoosen == None:
                del firstshape[1]
            print('firstshape', firstshape)
            print('secondshape', secondshape)
            print('radioval', radioval)
            if radioval == 'optionlibre' and valx2 != None and valy2 != None:

                lst = []
                for j in zip(valy2, valx2):
                    lst.append(j)
                s = -1
                m = ''
                for i in range(len(lst)):
                    if lst[i][0][-2].isdigit() == 1:
                        m = lst[i][0][-2]
                        m = 'T' + m
                    elif lst[i][0][-1].isdigit() == 1:
                        m = lst[i][0][-1]
                        m = 'T' + m
                    s += 1
                    a = df[lst[i][0]]
                    b = df[lst[i][1]]
                    if nclick > 0:
                        if axisdrop == lst[i][1]:
                            p = []
                            c = []
                            for t in df[lst[i][0]]:
                                if shift_x == None:
                                    raise PreventUpdate
                                else:
                                    t += float(shift_x)
                                    p.append(t)
                            df[lst[i][0]] = pd.DataFrame(p)
                            a = df[lst[i][0]]
                            df.to_excel("appending.xlsx")
                            for y in df[axisdrop]:
                                if shift_y == None:
                                    raise PreventUpdate
                                else:
                                    y += float(shift_y)
                                    c.append(y)
                            c.append(axisdrop)
                            df[axisdrop] = pd.DataFrame(c)
                            b = df[axisdrop]
                            df.to_excel("appending.xlsx")
                    print('valy2 neymis bakalim', valy2)
                    print('valx2 neymis bakalim', valx2)
                    for j in range(len(valy2)):
                        for k in range(len(valx2)):
                            a = df[valy2[j]]
                            b = df[valx2[k]]
                            fig.add_trace(go.Scatter(x=a, y=b, mode=radio, name="{}/{}".format(valy2[j], valx2[k])))
                            a = []
                            if nc > 0:
                                a = controlShape()
                            # if cleanclick > 0 :
                            #     a = []
                            fig.update_xaxes(
                                tickangle=90,
                                title_text='' if g1 == [] else g1[-1],
                                title_font={"size": 20},
                                title_standoff=25),

                            fig.update_yaxes(
                                title_text='' if g2 == [] else g2[-1],
                                title_standoff=25),
                            fig.update_layout(
                                title_text=head[-1] if len(head) > 0 else "{}/{}".format(valx2[0], valy2[0]),
                                autosize=True,
                                width=slidewidth,

                                shapes=a if (nc > cleanclick) else [],
                                height=slideheight,
                                margin=dict(
                                    l=50,
                                    r=50,
                                    b=50,
                                    t=50,
                                    pad=4
                                ),
                                # hovermode='x unified',
                                uirevision=valy2[0], ),
                            fig.add_annotation(text=note[-1] if len(note) > 0 else '',
                                               xref="paper", yref="paper",
                                               x=0, y=0.7, showarrow=False)

                    return fig

            if radioval == 'choosevalue' and len(valxsecond) > 0 and len(valysecond) > 0:
                lst = []
                for j in zip(valysecond, valxsecond):
                    lst.append(j)
                print('lst', lst)
                s = -1
                m = ''
                for i in range(len(lst)):
                    if lst[i][0][-2].isdigit() == 1:
                        m = lst[i][0][-2]
                        m = 'T' + m
                    elif lst[i][0][-1].isdigit() == 1:
                        m = lst[i][0][-1]
                        m = 'T' + m
                    s += 1
                    a = df[lst[i][0]]
                    b = df[lst[i][1]]
                    if nclick > 0:
                        if axisdrop == lst[i][1]:
                            p = []
                            c = []
                            for t in df[lst[i][0]]:
                                if shift_x == None:
                                    raise PreventUpdate
                                else:
                                    t += float(shift_x)
                                    p.append(t)
                            df[lst[i][0]] = pd.DataFrame(p)
                            a = df[lst[i][0]]
                            df.to_excel("appending.xlsx")
                            for y in df[axisdrop]:
                                if shift_y == None:
                                    raise PreventUpdate
                                else:
                                    y += float(shift_y)
                                    c.append(y)
                            c.append(axisdrop)
                            df[axisdrop] = pd.DataFrame(c)
                            b = df[axisdrop]
                            df.to_excel("appending.xlsx")

                    fig.add_trace(go.Scatter(x=a, y=b, mode=radio, name="{}/{}".format(valxsecond[s], valysecond[s])))

                    def controlShape():
                        pathline = ''
                        pathline2 = ''
                        minValfirst = 0
                        minValsecond = 0
                        val = 0
                        if firstchoosen != None and secondchoosen != None:
                            if len(firstshape) == 2 and leftfirstval != None and leftsecondval != None:
                                if int(firstshape[1]) > int(firstshape[0]):
                                    pathline = ''
                                    rangeshape = range(int(firstshape[0]), int(firstshape[1]))
                                    print('df[lst[i][0]][0]', df[lst[i][0]][0])
                                    if ':' or '-' in df[lst[i][0]][0]:
                                        print('burda miyim1')
                                        for k in rangeshape:
                                            if k == rangeshape[0]:
                                                if lst[i][1] == firstchoosen:
                                                    val = i
                                                print('df[lst[i][0]][k]', df[lst[i][0]][k])
                                                pathline += 'M ' + str(int(df[lst[val][0]][k])) + ', ' + str(
                                                    minValsecond) + ' L' + \
                                                            str(int(df[lst[val][0]][k])) + ', ' + str(
                                                    df[firstchoosen][k]) + ' '

                                            elif k != rangeshape[0] and k != rangeshape[-1]:
                                                if lst[i][1] == firstchoosen:
                                                    val = i
                                                pathline += ' L' + str(int(df[lst[val][0]][k])) + ', ' + str(
                                                    df[firstchoosen][k])
                                        pathline += ' L' + str(int(df[lst[val][0]][k])) + ', ' + str(minValsecond)
                                        pathline += ' Z'
                                    else:
                                        print('yoksa burda miyim')
                                        for k in rangeshape:
                                            if k == rangeshape[0]:
                                                pathline += 'M ' + str(int(df[lst[val][0]][k])) + ', ' + str(
                                                    minValfirst) + ' L' + \
                                                            str(int(df[lst[val][0]][k])) + ', ' + str(
                                                    df[firstchoosen][k]) + ' '

                                            elif k != rangeshape[0] and k != rangeshape[-1]:
                                                pathline += ' L' + str(int(df[lst[val][0]][k])) + ', ' + str(
                                                    df[firstchoosen][k])
                                        pathline += ' L' + str(int(df[lst[val][0]][k])) + ', ' + str(minValfirst)
                                        pathline += ' Z'
                                    print('pathline ==========>', pathline)

                            if len(secondshape) == 2 and rightsecondval != None and rightfirstval != None:
                                if int(secondshape[1]) > int(secondshape[0]):
                                    rangeshape = range(int(secondshape[0]), int(secondshape[1]))
                                    if ':' or '-' in df[lst[i][0]][0]:
                                        for k in rangeshape:
                                            if k == rangeshape[0]:
                                                print('df[lst[i][0]][k] 2. section', df[lst[i][0]])
                                                pathline2 += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(
                                                    minValsecond) + ' L' + \
                                                             str(df[lst[i][0]][k]) + ', ' + str(
                                                    df[secondchoosen][k]) + ' '

                                            elif k != rangeshape[0] and k != rangeshape[-1]:
                                                pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(
                                                    df[secondchoosen][k])
                                        pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond)
                                        pathline2 += ' Z'
                                    else:
                                        for k in rangeshape:
                                            if k == rangeshape[0]:
                                                pathline2 += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                                    minValsecond) + ' L' + \
                                                             str(int(df[lst[i][0]][k])) + ', ' + str(
                                                    df[secondchoosen][k]) + ' '

                                            elif k != rangeshape[0] and k != rangeshape[-1]:
                                                pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                                    df[secondchoosen][k])
                                        pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(minValsecond)
                                        pathline2 += ' Z'
                                    print('pathline222 ==========>', pathline2)
                            return [dict(
                                type="path",
                                path=pathline,
                                layer='below',
                                fillcolor="#5083C7",
                                opacity=0.8,
                                line_color="#8896BF",
                            ), dict(
                                type="path",
                                path=pathline2,
                                layer='below',
                                fillcolor="#B0384A",
                                opacity=0.8,
                                line_color="#B36873",
                            )]

                        if firstchoosen != None and secondchoosen == None:
                            if len(firstshape) == 2:
                                if int(firstshape[1]) > int(firstshape[0]):
                                    pathline = ''
                                    rangeshape = range(int(firstshape[0]), int(firstshape[1]))
                                    if ':' or '-' in df[lst[i][0]][0]:
                                        for k in rangeshape:
                                            if k == rangeshape[0]:
                                                pathline += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(
                                                    minValfirst) + ' L' + str(
                                                    df[lst[i][0]][k]) + ', ' + str(
                                                    df[firstchoosen][k]) + ' '

                                            elif k != rangeshape[0] and k != rangeshape[-1]:
                                                pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(
                                                    df[firstchoosen][k])
                                        pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValfirst)
                                        pathline += ' Z'
                                    else:
                                        for k in rangeshape:

                                            if k == rangeshape[0]:
                                                pathline += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                                    minValfirst) + ' L' + str(
                                                    int(df[lst[i][0]][k])) + ', ' + str(
                                                    df[firstchoosen][k]) + ' '

                                            elif k != rangeshape[0] and k != rangeshape[-1]:
                                                pathline += ' L' + str(int(a[k])) + ', ' + str(df[firstchoosen][k])
                                        pathline += ' L' + str(int(a[k])) + ', ' + str(minValfirst)
                                        pathline += ' Z'

                                    return [dict(
                                        type="path",
                                        path=pathline,
                                        layer='below',
                                        fillcolor="#5083C7",
                                        opacity=0.8,
                                        line_color="#8896BF",
                                    )]

                                if int(firstshape[0]) > int(firstshape[1]):
                                    rangeshape = range(int(firstshape[1]), int(firstshape[0]))
                                    if ':' or '-' in df[lst[i][0]][0]:
                                        for k in rangeshape:
                                            if k == rangeshape[0]:
                                                pathline += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(
                                                    minValsecond) + ' L' + str(
                                                    df[lst[i][0]][k]) + ', ' + str(
                                                    df[firstchoosen][k]) + ' '

                                            elif k != rangeshape[0] and k != rangeshape[-1]:
                                                pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(
                                                    df[firstchoosen][k])
                                        pathline += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond)
                                        pathline += ' Z'
                                    else:
                                        for k in rangeshape:

                                            if k == rangeshape[0]:
                                                pathline += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                                    minValsecond) + ' L' + str(
                                                    int(df[lst[i][0]][k])) + ', ' + str(
                                                    df[firstchoosen][k]) + ' '

                                            elif k != rangeshape[0] and k != rangeshape[-1]:
                                                pathline += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                                    df[firstchoosen][k])
                                        pathline += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(minValsecond)
                                        pathline += ' Z'

                                    return [dict(
                                        type="path",
                                        path=pathline,
                                        layer='below',
                                        fillcolor="#5083C7",
                                        opacity=0.8,
                                        line_color="#8896BF",
                                    )]

                        if secondchoosen != None and firstchoosen == None:
                            if len(secondshape) == 2 and rightsecondval != None and rightfirstval != None:
                                if int(secondshape[1]) > int(secondshape[0]):
                                    rangeshape = range(int(secondshape[0]), int(secondshape[1]))
                                    if ':' or '-' in df[lst[i][0]][0]:
                                        for k in rangeshape:
                                            if k == rangeshape[0]:
                                                pathline2 += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(
                                                    minValsecond) + ' L' + str(
                                                    df[lst[i][0]][k]) + ', ' + str(
                                                    df[secondchoosen][k]) + ' '

                                            elif k != rangeshape[0] and k != rangeshape[-1]:
                                                pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(
                                                    df[secondchoosen][k])
                                        pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond)
                                        pathline2 += ' Z'
                                    else:
                                        for k in rangeshape:

                                            if k == rangeshape[0]:
                                                pathline2 += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                                    minValsecond) + ' L' + str(
                                                    int(df[lst[i][0]][k])) + ', ' + str(
                                                    df[secondchoosen][k]) + ' '

                                            elif k != rangeshape[0] and k != rangeshape[-1]:
                                                pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                                    df[secondchoosen][k])
                                        pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(minValsecond)
                                        pathline2 += ' Z'

                                    return [dict(
                                        type="path",
                                        path=pathline2,
                                        layer='below',
                                        fillcolor="#5083C7",
                                        opacity=0.8,
                                        line_color="#8896BF",
                                    )]

                                if int(secondshape[0]) > int(secondshape[1]):
                                    rangeshape = range(int(secondshape[1]), int(secondshape[0]))
                                    for k in rangeshape:
                                        if k == rangeshape[0]:
                                            pathline2 += 'M ' + str(df[lst[i][0]][k]) + ', ' + str(
                                                minValsecond) + ' L' + str(
                                                df[lst[i][0]][k]) + ', ' + str(
                                                df[secondchoosen][k]) + ' '

                                        elif k != rangeshape[0] and k != rangeshape[-1]:
                                            pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(df[secondchoosen][k])
                                    pathline2 += ' L' + str(df[lst[i][0]][k]) + ', ' + str(minValsecond)
                                    pathline2 += ' Z'
                                else:
                                    rangeshape = range(int(secondshape[1]), int(secondshape[0]))
                                    for k in rangeshape:

                                        if k == rangeshape[0]:
                                            pathline2 += 'M ' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                                minValsecond) + ' L' + str(
                                                int(df[lst[i][0]][k])) + ', ' + str(
                                                df[secondchoosen][k]) + ' '

                                        elif k != rangeshape[0] and k != rangeshape[-1]:
                                            pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(
                                                df[secondchoosen][k])
                                    pathline2 += ' L' + str(int(df[lst[i][0]][k])) + ', ' + str(minValsecond)
                                    pathline2 += ' Z'

                                return [dict(
                                    type="path",
                                    path=pathline2,
                                    layer='below',
                                    fillcolor="#5083C7",
                                    opacity=0.8,
                                    line_color="#8896BF",
                                )]

                    a = []
                    if nc > 0:
                        a = controlShape()
                    fig.update_xaxes(
                        tickangle=90,
                        title_text='' if g1 == [] else g1[-1],
                        title_font={"size": 20},
                        title_standoff=25),

                    fig.update_yaxes(
                        title_text='' if g2 == [] else g2[-1],
                        title_standoff=25),
                    fig.update_shapes(yref='y'),
                    fig.update_layout(
                        title_text=head[-1] if len(head) > 0 else "{}/{}".format(valxsecond[0], valysecond[0]),
                        autosize=True,
                        width=slidewidth,
                        shapes=a if (nc > cleanclick) else [],
                        height=slideheight,
                        margin=dict(
                            l=50,
                            r=50,
                            b=50,
                            t=50,
                            pad=4
                        ),
                        yaxis=dict(
                            title='' if g2 == [] else g2[-1],
                            titlefont=dict(
                                color="#1f77b4"
                            ),
                            tickfont=dict(
                                color="#1f77b4"
                            )
                        ),
                        # yaxis2=dict(
                        #     title='' if g2 == [] else g2[-1],
                        #     titlefont=dict(
                        #         color="#d62728"
                        #     ),
                        #     tickfont=dict(
                        #         color="#d62728"
                        #     ),
                        #     anchor="x",
                        #     overlaying="y",
                        #     side="right"),
                        # hovermode='x unified',
                        uirevision=valysecond[0], ),
                    fig.add_annotation(text=note[-1] if len(note) > 0 else '',
                                       xref="paper", yref="paper",
                                       x=0, y=0.7, showarrow=False)

                return fig


            else:
                return no_update
        else:
            return no_update


@app.callback(
    [Output('pointLeftFirst', 'children'),
     Output('pointLeftSecond', 'children')],
    [Input('graph', 'clickData'),
     Input('firstChoosenValue', 'value'), ],
    [State('leftSideChecklistValueHidden', 'children'),
     State('pointLeftFirst', 'children'),
     State('pointLeftSecond', 'children'),
     State('shift_x_axis', 'value'),
     State('retrieve', 'children'),
     ]
)
def valint(clickData, firstchoosen, value, leftchild, rightchild, shift_x, retrieve):
    if value is [] or value is None or clickData == None or clickData == [] or firstchoosen == None or retrieve == None or retrieve == []:
        raise PreventUpdate

    spaceList1 = []
    zero = 0
    spaceList2 = []
    if len(retrieve) > 0:
        df = pd.read_excel('appending.xlsx')
        df['index'] = df.index
        for i in range(len(value)):
            spaceList1.append(zero)
            zero += 1
            spaceList2.append(value[i])
        zippedval = [i for i in list(zip(spaceList1, spaceList2))]
        curvenumber = clickData['points'][0]['curveNumber']
        for k in zippedval:
            if k[1] == firstchoosen:
                if k[0] == curvenumber:
                    x_val = clickData['points'][0]['x']
                    if 'date' in df.columns:
                        dff = df[df['date'] == x_val]
                    else:
                        a = ''
                        for v in df.columns:
                            if 'Temps' in v:
                                a += v
                                dff = df[df[v] == x_val]
                                if shift_x != 0:
                                    x_val -= shift_x
                                    dff = df[df[v] == x_val]

                    a = []
                    a.append(dff[firstchoosen].index)
                    for i in range(len(a)):
                        for j in a:
                            leftchild.append(j[i])

                    if len(leftchild) > 2:
                        leftchild.pop(0)
                    print('left2', leftchild)
                    return (leftchild, leftchild)
                else:
                    return (no_update, no_update)
            # else : return(no_update,no_update)
    else:
        return (no_update, no_update)

    # return left


@app.callback([Output('leftIntegralFirst', 'value'), Output('leftIntegralSecond', 'value')],
              [Input('pointLeftFirst', 'children'), Input('pointLeftSecond', 'children')],
              [State('firstChoosenValue', 'value')], )
def display_hover_data(leftchild, rightchild, firstchoosen):
    if leftchild == None or rightchild == None or leftchild == [] or rightchild == []:
        raise PreventUpdate

    minchild = 0
    maxchild = 0
    if len(leftchild) == 2:
        for i in range(len(leftchild)):
            if leftchild[0] < leftchild[1]:
                minchild = leftchild[0]
                maxchild = leftchild[1]
            else:
                minchild = leftchild[1]
                maxchild = leftchild[0]
    else:
        minchild = leftchild[0]
        maxchild = leftchild[0]

    if firstchoosen != '' and len(leftchild) == 2:
        return ('T ' + str(minchild), 'T ' + str(maxchild))
    else:
        return (no_update, no_update)


@app.callback(
    [Output('pointLeftFirstTab4', 'children'),
     Output('pointLeftSecondTab4', 'children')],
    [Input('graph4', 'clickData'),
     Input('radiographtab4hidden', 'children'),
     Input('firstChoosenValueTab4', 'value'),
     # Input('shiftaxisdroptab4hidden', 'children'),
     ],  # describe variable of shift
    [State('tab4hiddenValuey_axissecond', 'children'),
     State('tab4hiddenValuex_axissecond', 'children'),
     State('tab2hiddenValuey_axis', 'children'),
     State('tab2hiddenValuex_axis', 'children'),
     State('pointLeftFirstTab4', 'children'),
     State('pointLeftSecondTab4', 'children'),
     State('retrieve', 'children'),
     State('shift_x_axistab4', 'value'),  # shifting value of x_axis
     State('output_s', 'children')  # it takes values of tabdropdowntop and topdropdowntoptab4
     ]
)
def valintTab4(clickData4, radioval, firstchoosen, valysecond, valxsecond, valy, valx, leftchild, rightchild, retrieve,
               shift_x, container):
    if clickData4 == None or clickData4 == [] or firstchoosen == None or retrieve == None or retrieve == []:
        raise PreventUpdate
    spaceList1 = []
    zero = 0
    spaceList2 = []
    if len(retrieve) > 0:
        df = pd.read_excel('appending.xlsx')
        df['index'] = df.index
        df.dropna(axis=0, inplace=True)
        for i in range(len(container)):
            spaceList1.append(zero)
            zero += 1
            spaceList2.append(container[i])
        zippedval = [i for i in list(zip(spaceList1, spaceList2))]
        curvenumber = clickData4['points'][0]['curveNumber']
        for k in zippedval:
            if k[1] == firstchoosen:
                if k[0] == curvenumber:
                    if radioval == 'choosevalue':
                        if firstchoosen[-1].isdigit() == 1:
                            if valxsecond != []:
                                t = valxsecond.index(firstchoosen)
                                m = valysecond[t]
                                x_val = clickData4['points'][0]['x']
                                print("xvalllllllll", x_val)
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[firstchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])

                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                            else:
                                m = firstchoosen[-1:]
                                m = 'T' + m
                                x_val = clickData4['points'][0]['x']
                                print("xvalllllllll", x_val)
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[firstchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])

                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                        elif firstchoosen[-2].isdigit() == 1:
                            if valxsecond != []:
                                t = valxsecond.index(firstchoosen)
                                m = valysecond[t]
                                x_val = clickData4['points'][0]['x']
                                print("xvalllllllll", x_val)
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[firstchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])

                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                            else:
                                m = firstchoosen[-2:]
                                m = 'T' + m
                                x_val = clickData4['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[firstchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])
                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                        else:
                            if valxsecond != []:
                                print('nedir simdi burdaki firstchoosen', firstchoosen)
                                t = valxsecond.index(firstchoosen)
                                print('nedir simdi burdaki firstchoosen', firstchoosen)
                                m = valysecond[t]
                                print('m ne ola ki', m)
                                x_val = clickData4['points'][0]['x']
                                print('x_val left', x_val)
                                dff = df[df[m] == x_val]
                                print('dffffffleft', dff)
                                a = []
                                a.append(dff[firstchoosen].index)
                                print('aaaaaaaleft', a)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])
                                        print("leftchild1left", leftchild)

                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                print("leftchild2left", leftchild)
                                return (leftchild, leftchild)
                            else:
                                return (no_update, no_update)

                    elif radioval == 'optionlibre':
                        if valx != []:
                            print('valxxsxhshxshxsh, ', valx)
                            print('df', df)
                            t = valx.index(firstchoosen)
                            m = valy[t]
                            print('mmmmmm', m)
                            x_val = clickData4['points'][0]['x']
                            print('x_val left first', x_val)
                            dff = df[df[m] == x_val]
                            print('df[m]', df[m])
                            print('dffffffleft', dff)
                            # if 'date' in df.columns:
                            #     dff = df[df['date'] == x_val]
                            # else:
                            #     a = ''
                            #     for v in df.columns:
                            #         if 'Temps' in v:
                            #             a += v
                            #             dff = df[df[v] == x_val]
                            #             if shift_x != 0:
                            #                 x_val -= shift_x
                            #                 dff = df[df[v] == x_val]
                            a = []
                            a.append(dff[valx].index)
                            print('aaaaaaaleft', a)
                            for i in range(len(a)):
                                for j in a:
                                    leftchild.append(j[i])
                                    print("leftchild1dsdsd", leftchild)

                            if len(leftchild) > 2:
                                leftchild.pop(0)
                            print("leftchild2sdsds", leftchild)

                            return (leftchild, leftchild)
                        else:
                            return (no_update, no_update)

                    else:
                        return (no_update, no_update)
                else:
                    return (no_update, no_update)

    else:
        return (no_update, no_update)


@app.callback([Output('leftIntegralFirstTab4', 'value'),
               Output('leftIntegralSecondTab4', 'value')],
              [Input('pointLeftFirstTab4', 'children'),
               Input('pointLeftSecondTab4', 'children'),
               Input('firstChoosenValueTab4', 'value'), ], )
def display_hover_dataTab4(leftchild1, rightchild, firstchoosen):
    if leftchild1 == None or firstchoosen == None or rightchild == None or leftchild1 == [] or rightchild == []:
        raise PreventUpdate

    if firstchoosen != '' and len(leftchild1) == 2:
        for i in range(len(leftchild1)):
            if leftchild1[0] < leftchild1[1]:
                minchild = leftchild1[0]
                maxchild = leftchild1[1]
                return 'T ' + str(minchild), 'T ' + str(maxchild)
            else:
                minchild = leftchild1[1]
                maxchild = leftchild1[0]
                return 'T ' + str(minchild), 'T ' + str(maxchild)
    else:
        return no_update, no_update


@app.callback(
    [Output('pointRightFirst', 'children'),
     Output('pointRightSecond', 'children')],
    [Input('graph', 'clickData'),
     Input('secondChoosenValue', 'value')],
    [State('leftSideChecklistValueHidden', 'children'),
     State('pointRightFirst', 'children'),
     State('pointRightSecond', 'children'),
     State('shift_x_axis', 'value'),
     State('retrieve', 'children')]
)
def valint2(clickData, secondchoosen, value, leftchild, rightchild, shift_x, retrieve):
    if value is [] or value is None or clickData == None or secondchoosen == None or retrieve == None or retrieve == []:
        raise PreventUpdate

    spaceList1 = []
    zero = 0
    spaceList2 = []
    if len(retrieve) > 0:
        df = pd.read_excel('appending.xlsx')
        df['index'] = df.index
        for i in range(len(value)):
            spaceList1.append(zero)
            zero += 1
            spaceList2.append(value[i])
        zippedval = [i for i in list(zip(spaceList1, spaceList2))]
        curvenumber = clickData['points'][0]['curveNumber']
        for k in zippedval:
            if k[1] == secondchoosen:
                if k[0] == curvenumber:
                    x_val = clickData['points'][0]['x']
                    if 'date' in df.columns:
                        dff = df[df['date'] == x_val]
                    else:
                        a = ''
                        for v in df.columns:
                            if 'Temps' in v:
                                a += v
                                dff = df[df[v] == x_val]
                                if shift_x != 0:
                                    x_val -= shift_x
                                    dff = df[df[v] == x_val]
                    a = []
                    a.append(dff[secondchoosen].index)
                    for i in range(len(a)):
                        for j in a:
                            leftchild.append(j[i])
                    if len(leftchild) > 2:
                        leftchild.pop(0)
                    return (leftchild, leftchild)
                else:
                    return (no_update, no_update)
            # else : return (no_update, no_update)
    else:
        return (no_update, no_update)


@app.callback(
    [Output('rightIntegralFirst', 'value'), Output('rightIntegralSecond', 'value')],
    [Input('pointRightFirst', 'children'), Input('pointRightSecond', 'children')],
    [State('secondChoosenValue', 'value')], )
def display_hover_data2(leftchild1, rightchild1, secondchoosen):
    if leftchild1 == None or rightchild1 == None or leftchild1 == [] or rightchild1 == [] or secondchoosen == None:
        raise PreventUpdate
    if secondchoosen != '' and len(leftchild1) == 2:
        for i in range(len(leftchild1)):
            if leftchild1[0] < leftchild1[1]:
                minchild = leftchild1[0]
                maxchild = leftchild1[1]
                return 'T ' + str(minchild), 'T ' + str(maxchild)
            else:
                minchild = leftchild1[1]
                maxchild = leftchild1[0]
                return 'T ' + str(minchild), 'T ' + str(maxchild)
    else:
        minchild = leftchild1[0]
        maxchild = leftchild1[0]

    if secondchoosen != '' and len(leftchild1) == 2:
        return 'T ' + str(minchild), 'T ' + str(maxchild)
    else:
        return (no_update, no_update)


@app.callback(
    [Output('pointRightFirstTab4', 'children'),
     Output('pointRightSecondTab4', 'children')],
    [Input('graph4', 'clickData'),
     Input('radiographtab4hidden', 'children'),
     Input('secondChoosenValueTab4', 'value'),
     ],
    [State('tab4hiddenValuey_axissecond', 'children'),
     State('tab4hiddenValuex_axissecond', 'children'),
     State('tab2hiddenValuey_axis', 'children'),
     State('tab2hiddenValuex_axis', 'children'),
     State('pointRightFirstTab4', 'children'),
     State('pointRightSecondTab4', 'children'),
     State('retrieve', 'children'),
     State('output_s', 'children'),
     State('shift_x_axistab4', 'value'), ]
)
def valintTab4_2(clickData, radioval, secondchoosen, valysecond, valxsecond, valy, valx, leftchild, rightchild,
                 retrieve, container, shift_x):
    if clickData == None or container is [] or container is None or secondchoosen == None or secondchoosen == [] or retrieve == None or retrieve == []:
        raise PreventUpdate

    spaceList1 = []
    zero = 0
    spaceList2 = []
    if len(retrieve) > 0:
        df = pd.read_excel('appending.xlsx')
        df['index'] = df.index
        for i in range(len(container)):
            spaceList1.append(zero)
            zero += 1
            spaceList2.append(container[i])
        zippedval = [i for i in list(zip(spaceList1, spaceList2))]
        curvenumber = clickData['points'][0]['curveNumber']
        for k in zippedval:
            if k[1] == secondchoosen:
                if k[0] == curvenumber:
                    if radioval == "choosevalue":
                        if secondchoosen[-1].isdigit() == 1:
                            print('valxsecond ne alaka anlamadim 1 ', valxsecond)
                            if valxsecond != []:
                                t = valxsecond.index(secondchoosen)
                                m = valysecond[t]
                                x_val = clickData['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[secondchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])
                                    if len(leftchild) > 2:
                                        leftchild.pop(0)
                                return (leftchild, leftchild)
                            else:
                                m = secondchoosen[-1:]
                                m = 'T' + m
                                x_val = clickData['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[secondchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])
                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                        elif secondchoosen[-2].isdigit() == 1:
                            print('valxsecond ne alaka anlamadim 2 ', valxsecond)
                            if valxsecond != []:
                                t = valxsecond.index(secondchoosen)
                                m = valysecond[t]
                                x_val = clickData['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[secondchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])
                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                            else:
                                m = secondchoosen[-2:]
                                m = 'T' + m
                                x_val = clickData['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[secondchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])
                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                        else:
                            if valxsecond != []:
                                print('valxsecond ne alaka anlamadim else', valxsecond)
                                print('secondchoosen nedir', secondchoosen)
                                t = valxsecond.index(secondchoosen)
                                m = valysecond[t]
                                print('buradaki m nedir karmasik oldu', m)
                                x_val = clickData['points'][0]['x']
                                dff = df[df[m] == x_val]
                                a = []
                                a.append(dff[secondchoosen].index)
                                for i in range(len(a)):
                                    for j in a:
                                        leftchild.append(j[i])

                                if len(leftchild) > 2:
                                    leftchild.pop(0)
                                return (leftchild, leftchild)
                            else:
                                return no_update, no_update
                    elif radioval == 'optionlibre':
                        if valx != []:
                            print('valxxxxxx', valx)
                            print('valyyyyyy', valy)
                            t = valx.index(secondchoosen)
                            m = valy[0]
                            print('mmmmmmmmmmmchange', m)
                            x_val = clickData['points'][0]['x']
                            dff = df[df[m] == x_val]
                            a = []
                            a.append(dff[secondchoosen].index)
                            print('aaaaaaaachange', a)
                            for i in range(len(a)):
                                for j in a:
                                    leftchild.append(j[i])

                            if len(leftchild) > 2:
                                leftchild.pop(0)
                            return (leftchild, leftchild)
                        else:
                            return (no_update, no_update)

                    else:
                        return (no_update, no_update)
                else:
                    return (no_update, no_update)
            # else:
            #     return (no_update, no_update)

    else:
        return (no_update, no_update)


@app.callback(
    [Output('rightIntegralFirstTab4', 'value'),
     Output('rightIntegralSecondTab4', 'value')],
    [Input('pointRightFirstTab4', 'children'),
     Input('pointRightSecondTab4', 'children'),
     Input('secondChoosenValueTab4', 'value'), ], )
def display_hover_data4(leftchild1, rightchild1, secondchoosen):
    if leftchild1 == None or rightchild1 == None or leftchild1 == [] or rightchild1 == [] or secondchoosen == None:
        raise PreventUpdate

    if secondchoosen != '' and len(leftchild1) == 2:
        for i in range(len(leftchild1)):
            if leftchild1[0] < leftchild1[1]:
                minchild = leftchild1[0]
                maxchild = leftchild1[1]
                return 'T ' + str(minchild), 'T ' + str(maxchild)
            else:
                minchild = leftchild1[1]
                maxchild = leftchild1[0]
                return 'T ' + str(minchild), 'T ' + str(maxchild)

    else:
        return no_update, no_update


@app.callback(Output('leftIntegral', 'value'),
              [Input('leftIntegralFirst', 'value'),
               Input('leftIntegralSecond', 'value'),
               Input('firstChoosenValue', 'value'),
               ], [State('retrieve', 'children')]
              )
def integralCalculation(st1left, st1right, valuechoosenleft, retrieve):
    if st1left == None or st1right == None or valuechoosenleft == None or valuechoosenleft == [] or retrieve == None or retrieve == []:
        raise PreventUpdate
    if st1left.startswith('T') == 1 and st1right.startswith('T') == 1:
        st1left = st1left[2:]
        st1right = st1right[2:]
    elif st1left.startswith('T') == 1 and st1right.isnumeric() == 1:
        st1left = st1left[2:]
        st1right = st1right
    elif st1left.isnumeric() == 1 and st1right.isnumeric() == 1:
        st1left = st1left
        st1right = st1right
    elif st1left.isnumeric() == 1 and st1right.startswith('T') == 1:
        st1left = st1left
        st1right = st1right[2:]
    if len(retrieve) > 0:
        if st1left != '' and st1right != '':
            df = pd.read_excel('appending.xlsx')
            df['index'] = df.index
            df = df.reindex(columns=sorted(df.columns, reverse=True))
            dff1 = df[(df[valuechoosenleft].index >= float(st1left)) & (df[valuechoosenleft].index <= float(st1right)) |
                      (df[valuechoosenleft].index >= float(st1right)) & (df[valuechoosenleft].index <= float(st1left))]
            for i in df.columns:
                if i.startswith('Temps'):
                    dff1 = dff1.groupby(i).mean()

            c = dff1[valuechoosenleft]
            area1 = abs(trapz((abs(c)), dx=1))

            return area1
        elif (st1left == '' and st1right != '') or (st1left != '' and st1right == ''):
            return 'total integration'
        elif (st1left == '' and st1right == '') and valuechoosenleft != '':
            return 'total integration'
        elif st1left != '' and st1right != '' and valuechoosenleft == '':
            return 'total integration'
    # return no_update


@app.callback(Output('leftIntegralTab4', 'value'),
              [Input('leftIntegralFirstTab4', 'value'),
               Input('leftIntegralSecondTab4', 'value'),
               Input('firstChoosenValueTab4', 'value'),
               ], [State('retrieve', 'children')]
              )
def integralCalculationtab4(st1left, st1right, valuechoosenleft, retrieve):
    if st1left == None or st1right == None or valuechoosenleft == None or valuechoosenleft == [] or retrieve == None or retrieve == []:
        raise PreventUpdate
    if st1left.startswith('T') == 1 and st1right.startswith('T') == 1:
        st1left = st1left[2:]
        st1right = st1right[2:]
    elif st1left.startswith('T') == 1 and st1right.isnumeric() == 1:
        st1left = st1left[2:]
        st1right = st1right
    elif st1left.isnumeric() == 1 and st1right.isnumeric() == 1:
        st1left = st1left
        st1right = st1right
    elif st1left.isnumeric() == 1 and st1right.startswith('T') == 1:
        st1left = st1left
        st1right = st1right[2:]
    if len(retrieve) > 0:
        if st1left != '' and st1right != '':
            df = pd.read_excel('appending.xlsx')
            df['index'] = df.index
            df = df.reindex(columns=sorted(df.columns, reverse=True))
            dff1 = df[(df[valuechoosenleft].index >= float(st1left)) & (df[valuechoosenleft].index <= float(st1right)) |
                      (df[valuechoosenleft].index >= float(st1right)) & (df[valuechoosenleft].index <= float(st1left))]
            for i in df.columns:
                if i.startswith('Temps'):
                    dff1 = dff1.groupby(i).mean()
            c = dff1[valuechoosenleft]
            area1 = abs(trapz(abs(c), dx=1))

            return area1
        elif (st1left == '' and st1right != '') or (st1left != '' and st1right == ''):
            return 'total integration'
        elif (st1left == '' and st1right == '') and valuechoosenleft != '':
            return 'total integration'
        elif st1left != '' and st1right != '' and valuechoosenleft == '':
            return 'total integration'
    # return no_update


@app.callback(Output('rightIntegral', 'value'),
              [Input('rightIntegralFirst', 'value'),
               Input('rightIntegralSecond', 'value'),
               Input('secondChoosenValue', 'value'),
               ], [State('retrieve', 'children')]
              )
def integralCalculation2(st2left, st2right, valuechoosenright, retrieve):
    if st2left == None or st2right == None or valuechoosenright == None or valuechoosenright == [] or retrieve == None or retrieve == []:
        raise PreventUpdate
    if st2left.startswith('T') == 1 and st2right.startswith('T') == 1:
        st2left = st2left[2:]
        st2right = st2right[2:]
    elif st2left.startswith('T') == 1 and st2right.isnumeric() == 1:
        st2left = st2left[2:]
        st2right = st2right
    elif st2left.isnumeric() == 1 and st2right.isnumeric() == 1:
        st2left = st2left
        st2right = st2right
    elif st2left.isnumeric() == 1 and st2right.startswith('T') == 1:
        st2left = st2left
        st2right = st2right[2:]
    if len(retrieve) > 0:
        if st2left != '' and st2right != '':
            df = pd.read_excel('appending.xlsx')
            df['index'] = df.index
            df = df.reindex(columns=sorted(df.columns, reverse=True))
            dff2 = df[
                (df[valuechoosenright].index >= float(st2left)) & (df[valuechoosenright].index <= float(st2right)) |
                (df[valuechoosenright].index >= float(st2right)) & (df[valuechoosenright].index <= float(st2left))]
            for i in df.columns:
                if i.startswith('Temps'):
                    dff2 = dff2.groupby(i).mean()

            f = dff2[valuechoosenright]
            area2 = abs(trapz(abs(f), dx=1))
            return area2
        elif (st2left == '' and st2right != '') or (st2left != '' and st2right == ''):
            return 'total integration'
        elif (st2left == '' and st2right == '') and valuechoosenright != '':
            return 'total integration'
        elif st2left != '' and st2right != '' and valuechoosenright == '':
            return 'total integration'


@app.callback(Output('rightIntegralTab4', 'value'),
              [Input('rightIntegralFirstTab4', 'value'),
               Input('rightIntegralSecondTab4', 'value'),
               Input('secondChoosenValueTab4', 'value'),
               ], [State('retrieve', 'children')]
              )
def integralCalculation4(st2left, st2right, valuechoosenright, retrieve):
    if st2left == None or st2right == None or valuechoosenright == None or valuechoosenright == [] or retrieve == None or retrieve == []:
        raise PreventUpdate
    if st2left.startswith('T') == 1 and st2right.startswith('T') == 1:
        st2left = st2left[2:]
        st2right = st2right[2:]
    elif st2left.startswith('T') == 1 and st2right.isnumeric() == 1:
        st2left = st2left[2:]
        st2right = st2right
    elif st2left.isnumeric() == 1 and st2right.isnumeric() == 1:
        st2left = st2left
        st2right = st2right
    elif st2left.isnumeric() == 1 and st2right.startswith('T') == 1:
        st2left = st2left
        st2right = st2right[2:]
    if len(retrieve) > 0:
        if st2left != '' and st2right != '':
            df = pd.read_excel('appending.xlsx')
            df['index'] = df.index
            df = df.reindex(columns=sorted(df.columns, reverse=True))
            dff2 = df[
                (df[valuechoosenright].index >= float(st2left)) & (df[valuechoosenright].index <= float(st2right)) |
                (df[valuechoosenright].index >= float(st2right)) & (df[valuechoosenright].index <= float(st2left))]
            for i in df.columns:
                if i.startswith('Temps'):
                    dff2 = dff2.groupby(i).mean()
            f = dff2[valuechoosenright]
            area2 = abs(trapz(abs(f), dx=1))
            return area2
        elif (st2left == '' and st2right != '') or (st2left != '' and st2right == ''):
            return 'total integration'
        elif (st2left == '' and st2right == '') and valuechoosenright != '':
            return 'total integration'
        elif st2left != '' and st2right != '' and valuechoosenright == '':
            return 'total integration'


@app.callback(Output('operation', 'value'),
              [Input('leftIntegral', 'value'),
               Input('rightIntegral', 'value'),
               Input('operateur', 'value')],
              )
def differanceintegration(value1, value2, ops):
    if value1 == None or value2 == None:
        raise PreventUpdate
    if ops == ['Plus']:
        return float(value1 + value2)
    elif ops == ['Moins']:
        return float(value1 - value2)
    elif ops == ['Multiplie']:
        return float(value1 * value2)
    elif ops == ['Division']:
        return float(value1 / value2)
    elif ops == []:
        return []


@app.callback(Output('operationTab4', 'value'),
              [Input('leftIntegralTab4', 'value'),
               Input('rightIntegralTab4', 'value'),
               Input('operateurTab4', 'value')],
              )
def differanceintegrationTab4(value1, value2, ops):
    if value1 == None or value2 == None:
        raise PreventUpdate
    if ops == ['Plus']:
        return float(value1 + value2)
    elif ops == ['Moins']:
        return float(value1 - value2)
    elif ops == ['Multiplie']:
        return float(value1 * value2)
    elif ops == ['Division']:
        return float(value1 / value2)
    elif ops == []:
        return []


@app.callback(Output('intersection', 'value'),
              [Input('hiddenDifferance', 'children'),
               Input('firstChoosenValue', 'value'),
               Input('secondChoosenValue', 'value'),
               Input('leftIntegralFirst', 'value'),
               Input('rightIntegralFirst', 'value'), ],
              [State('intersection', 'value'), State('retrieve', 'children'),
               ]
              )
def differanceCalculation(hiddendif, valuechoosenleft, valuechoosenright, leftfirst, rightfirst, diff, retrieve):
    if hiddendif == None or hiddendif == [] or retrieve == None or retrieve == []:
        raise PreventUpdate

    # (len(hiddendif)>=2 and len(valuechoosenright)==1) or (len(hiddendif)>=2 and len(valuechoosenleft)==1) or
    if (len(hiddendif) >= 2):
        a = 0
        b = 0
        for i in range(len(hiddendif)):
            if hiddendif[0] < hiddendif[1]:
                a = hiddendif[0]
                b = hiddendif[1]
            else:
                a = hiddendif[1]
                b = hiddendif[0]
        differance = []
        if len(retrieve) > 0 and valuechoosenright != None and valuechoosenleft != None and leftfirst != None and rightfirst != None:

            df = pd.read_excel('appending.xlsx')
            df['index'] = df.index
            df = df.reindex(columns=sorted(df.columns, reverse=True))
            dff = df[(df[valuechoosenright].index >= float(a)) & (df[valuechoosenright].index <= float(b)) |
                     (df[valuechoosenright].index >= float(b)) & (df[valuechoosenright].index <= float(a))]
            l = dff[valuechoosenright]

            dff2 = df[(df[valuechoosenleft].index >= float(a)) & (df[valuechoosenleft].index <= float(b)) |
                      (df[valuechoosenleft].index >= float(b)) & (df[valuechoosenleft].index <= float(a))]
            r = dff2[valuechoosenleft]
            tt = []
            yy = []
            for i in l:
                tt.append(i)
            for i in r:
                yy.append(i)
            for i in range(len(tt)):
                if tt[i] <= yy[i]:
                    differance.append(tt[i])
                if yy[i] < tt[i]:
                    differance.append(yy[i])
            diff = (abs(trapz(differance, dx=1)))
            return diff
        else:
            return ['intersection']


@app.callback(Output('intersectionTab4', 'value'),
              [Input('pointLeftFirstTab4', 'children'),
               Input('pointRightFirstTab4', 'children'),
               Input('firstChoosenValueTab4', 'value'),
               Input('secondChoosenValueTab4', 'value'),
               Input('leftIntegralFirstTab4', 'value'),
               Input('rightIntegralFirstTab4', 'value'), ],
              [State('intersectionTab4', 'value'), State('retrieve', 'children'),
               ]
              )
def differanceCalculation4(firstshape, secondshape, valuechoosenleft, valuechoosenright, leftfirst, rightfirst, diff,
                           retrieve):
    if retrieve == None or retrieve == []:
        raise PreventUpdate
    differance = []
    if len(firstshape) == 2 and len(secondshape) == 2:
        a = int(firstshape[0])
        c = int(secondshape[0])
        b = int(firstshape[1])
        d = int(secondshape[1])
        if set(range(a, b)).issuperset(set(range(c, d))) == 1:
            differance.append(c)
            differance.append(d)
        elif set(range(c, d)).issuperset(set(range(a, b))) == 1:
            differance.append(a)
            differance.append(b)
        elif len(set(range(a, b)).intersection(set(range(c, d)))) >= 1 or len(
                set(range(c, d)).intersection(set(range(a, b)))) >= 1:
            if a <= c:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(b)
                differance.append(b)
            if a >= c:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(a)
                differance.append(a)
            if b <= d:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(c)
                differance.append(c)
            if b >= d:
                if len(differance) == 2:
                    differance.pop(0)
                    differance.append(d)
                differance.append(d)
        else:
            return ['intersection']
        differancelast = []
        if len(retrieve) > 0 and valuechoosenright != None and valuechoosenleft != None and leftfirst != None and rightfirst != None:
            first = differance[0]
            second = differance[1]
            df = pd.read_excel('appending.xlsx')
            df['index'] = df.index
            df = df.reindex(columns=sorted(df.columns, reverse=True))
            dff = df[(df[valuechoosenright].index >= float(first)) & (df[valuechoosenright].index <= float(second)) |
                     (df[valuechoosenright].index >= float(second)) & (df[valuechoosenright].index <= float(first))]
            l = dff[valuechoosenright]

            dff2 = df[(df[valuechoosenleft].index >= float(first)) & (df[valuechoosenleft].index <= float(second)) |
                      (df[valuechoosenleft].index >= float(second)) & (df[valuechoosenleft].index <= float(first))]
            r = dff2[valuechoosenleft]
            tt = []
            yy = []
            for i in l:
                tt.append(i)
            for i in r:
                yy.append(i)
            for i in range(len(tt)):
                if tt[i] <= yy[i]:
                    differancelast.append(tt[i])
                if yy[i] < tt[i]:
                    differancelast.append(yy[i])
            diff = (abs(trapz(differancelast, dx=1)))
            return diff


@app.callback(Output('writeexcelhidden', 'children'),
              [Input('write_excel', 'n_clicks')],
              [State('firstChoosenValue', 'value'),
               State('leftIntegralFirst', 'value'),
               State('leftIntegralSecond', 'value'),
               State('leftIntegral', 'value'),
               State('secondChoosenValue', 'value'),
               State('rightIntegralFirst', 'value'),
               State('rightIntegralSecond', 'value'),
               State('rightIntegral', 'value'),
               State('operation', 'value'),
               State('intersection', 'value'),
               ],
              )
def write_excel(nc, a, b, c, d, e, f, g, h, i, j):
    if nc > 0:
        now = datetime.datetime.now()
        if i == []:
            i = None
        if j == ['intersection']:
            j = None
        x = (now, a, b, c, d, e, f, g, h, i, j)

        if x != None: return x


@app.callback(Output('writeexcelhiddenTab4', 'children'),
              [Input('write_excelTab4', 'n_clicks')],
              [State('firstChoosenValueTab4', 'value'),
               State('leftIntegralFirstTab4', 'value'),
               State('leftIntegralSecondTab4', 'value'),
               State('leftIntegralTab4', 'value'),
               State('secondChoosenValueTab4', 'value'),
               State('rightIntegralFirstTab4', 'value'),
               State('rightIntegralSecondTab4', 'value'),
               State('rightIntegralTab4', 'value'),
               State('operationTab4', 'value'),
               State('intersectionTab4', 'value'),
               ],
              )
def write_excelTab4(nc, a, b, c, d, e, f, g, h, i, j):
    if nc > 0:
        now = datetime.datetime.now()
        if i == []:
            i = None
        if j == ['intersection']:
            j = None
        x = (now, a, b, c, d, e, f, g, h, i, j)

        if x != None: return x


@app.callback(Output('hiddenrecord3', 'children'),
              [Input('writeexcelhidden', 'children'), Input('writeexcelhiddenTab4', 'children')],
              )
def pasfunc(hiddenvalchild, hiddenvalchild4):
    if hiddenvalchild == None and hiddenvalchild4 == None:
        raise PreventUpdate
    if hiddenvalchild != None:
        return hiddenvalchild
    if hiddenvalchild4 != None:
        return hiddenvalchild4


@app.callback(Output('hiddenrecord4', 'children'),
              [Input('hiddenrecord3', 'children')],
              State('hiddenrecord4', 'children'), )
def lastfunc(hiddenvalchild, lastvalchild):
    lastvalchild = hiddenvalchild + lastvalchild
    return lastvalchild


@app.callback(Output('hiddenrecord1', 'children'),
              [Input('hiddenrecord4', 'children')],
              )
def exportdata(valueparse):
    a_parse = []
    t_parse = []
    for i in valueparse:
        if i == None:
            a_parse.append('None')
        else:
            a_parse.append(i)
        if len(a_parse) % 11 == 0:
            t_parse.append(a_parse)
            a_parse = []
    t_parse.insert(0, ['time', 'firstChoosenValue', 'leftIntegralFirst', 'leftIntegralSecond', 'leftIntegral',
                       'secondChoosenValue',
                       'rightIntegralFirst', 'rightIntegralSecond', 'rightIntegral', 'operation', 'intersection'])

    df = pd.DataFrame(t_parse)
    df.to_excel('new_fichier.xlsx')


@app.server.route("/download_excel/")
def download_excel():
    # Create DF
    dff = pd.read_excel("new_fichier.xlsx")
    # Convert DF
    buf = io.BytesIO()
    excel_writer = pd.ExcelWriter(buf, engine="xlsxwriter")
    dff.to_excel(excel_writer, sheet_name="sheet1")
    excel_writer.save()
    excel_data = buf.getvalue()
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        attachment_filename="LERMAB.xlsx",
        as_attachment=True,
        cache_timeout=0
    )


@app.callback(Output('Dbdesign', 'children'),
              [Input('tabs-with-classes', 'value')],
              )
def DBcall(tab):
    if tab == 'tab-3':
        datalist = html.Div([html.Div([html.Div([dbc.Button("Database Activate", id="activatedb", n_clicks=0,
                                                            color="success", size='lg',
                                                            ),
                                                 dbc.Input(id='db_Ip',
                                                           type="text",
                                                           debounce=True,
                                                           min=-10000, max=10000, step=1,
                                                           bs_size="mr",
                                                           style={'width': '11rem', "marginTop": "1.5rem"},
                                                           autoFocus=True,
                                                           placeholder="Enter your IP number ..."),
                                                 dbc.Input(id='db_name',
                                                           type="text",
                                                           debounce=True,
                                                           min=-10000, max=10000, step=1,
                                                           bs_size="mr",
                                                           style={'width': '11rem', "marginTop": "1.5rem"},
                                                           autoFocus=True,
                                                           placeholder="Enter your DB name...")
                                                 ], className="aadb"),
                                       html.Div([
                                           dcc.Dropdown(id='dbvalchoosen',
                                                        # options=[{'label': i, 'value': i}
                                                        #          for i in df.columns],
                                                        multi=False,
                                                        style={"cursor": "pointer", 'marginTop': '5px'},
                                                        className='stockSelectorClass3',
                                                        clearable=True,
                                                        placeholder='Select sent or received ...',

                                                        ),

                                           dcc.Dropdown(id='dbvalname',
                                                        # options=[{'label': i, 'value': i}
                                                        #          for i in df.columns],
                                                        multi=True,
                                                        style={"cursor": "pointer", 'marginTop': '13px'},
                                                        className='stockSelectorClass3',
                                                        clearable=True,
                                                        placeholder='Select your parameters...',
                                                        ),

                                           dcc.Dropdown(id='dbvaldate',
                                                        # options=[{'label': i, 'value': i}
                                                        #          for i in df.columns],
                                                        multi=True,
                                                        style={"cursor": "pointer", 'marginTop': '13px'},
                                                        className='stockSelectorClass3',
                                                        clearable=False,
                                                        placeholder='Select your parameters...',
                                                        ), ], className='aadb'), ], className='abcdb'),
                             dcc.Store(id='memory-output'),
                             html.Div([html.Div(dcc.Graph(id="getdbgraph",
                                                          config={'displayModeBar': True,
                                                                  'scrollZoom': True,
                                                                  'modeBarButtonsToAdd': [
                                                                      'drawline',
                                                                      'drawrect',
                                                                      'drawopenpath',
                                                                      'select2d',
                                                                      'eraseshape',
                                                                  ]},
                                                          style={'marginTop': 20, },
                                                          figure={
                                                              'layout': {'legend': {'tracegroupgap': 0},

                                                                         }
                                                          }

                                                          ), ),
                                       html.Div(daq.Slider(id="sliderHeightdb",
                                                           max=2100,
                                                           min=400,
                                                           value=530,
                                                           step=100,
                                                           size=400,
                                                           vertical=True,
                                                           updatemode='drag'), style={'margin': '20px'})],
                                      className='abcdb'),
                             html.Div([daq.Slider(id="sliderWidthdb",
                                                  max=2000,
                                                  min=600,
                                                  value=1000,
                                                  step=100,
                                                  size=750,
                                                  updatemode='drag'),
                                       html.Div(id='output-data-upload', children=[])]),
                             html.Div(dash_table.DataTable(id="getdbtable",
                                                           editable=True,
                                                           page_size=50,
                                                           style_table={'height': '500px', 'overflowY': 'auto',
                                                                        'width': '98%'},
                                                           style_cell={
                                                               'overflow': 'hidden',
                                                               'textOverflow': 'ellipsis',
                                                               'maxWidth': 0,
                                                               'fontSize': '1rem',
                                                               'TextAlign': 'center',
                                                           },
                                                           fixed_rows={'headers': True},

                                                           # style_cell_conditional=[
                                                           # {'if': {'column_id': 'date'},
                                                           #  'width': '15%'}

                                                           style_header={
                                                               'backgroundColor': 'rgb(230, 230, 230)',
                                                               'fontWeight': 'bold'
                                                           },
                                                           filter_action="native",
                                                           sort_action="native",
                                                           sort_mode="multi",
                                                           column_selectable="single",
                                                           # row_selectable="multi",
                                                           # row_deletable=True,
                                                           selected_columns=[],
                                                           selected_rows=[],
                                                           page_action="native",
                                                           page_current=0,
                                                           export_format='xlsx',
                                                           export_headers='display',
                                                           merge_duplicate_headers=True)),
                             html.Div(id="hiddendb1", children=[], style={'display': 'None'}),
                             html.Div(id="hiddendb2", style={'display': 'None'}),
                             html.Div(id="hiddendb3", children=[], style={'display': 'None'})

                             ], ),

        return datalist


# @app.callback(Output('dbvalchoosen', 'options'),
#               [Input('activatedb', 'n_clicks')],
#                [State('db_Ip', 'value'),
#                State('db_name', 'value')],
#               )
# def connectiondb(button,ipval,db_name):
#     if ipval == None and db_name == None:
#         raise PreventUpdate
#
#     if button > 0:
#
#         server = SSHTunnelForwarder(
#                 ("193.54.2.211", 22),
#                 ssh_username='soudani',
#                 ssh_password="univ484067152",
#                 remote_bind_address=("193.54.2.211", 3306))
#
#         server.start()
#
#         try:
#             conn = mariadb.connect(
#                     user="dashapp",
#                     password="dashapp",
#                     host="193.54.2.211",
#                     port=3306,
#                     database="rcckn"
#             )
#
#         except mariadb.Error as e:
#             print(f"Error connecting to MariaDB Platform: {e}")
#             sys.exit(1)
#             # Get Cursor
#         cur = conn.cursor()
#             # cur.execute("SELECT * FROM received_variablevalues WHERE LOCAL_TIMESTAMP <'2020-07-22 18:11:24'")
#         b = "select table_name from information_schema.tables where TABLE_SCHEMA='rcckn'"
#             # a = "SELECT DISTINCT VARIABLE_NAME FROM received_variablevalues "
#
#         cur.execute(b)
#         t = cur.fetchall()
#         df = pd.DataFrame(t)
#         m = []
#         for i in t:
#             m.append(i[0])
#         return [{'label': i, 'value': i} for i in m if i != 'app_variablerequest' and i != 'send_controlvalues' and
#                 i != 'received_ack' and i != 'send_vw_variablerequestdestination' and i != 'flyway_schema_history'
#                 and i != 'app_vw_messaging_followup' and i != 'received_variablerequest' and i != 'received_controlvalues'
#                 and i != 'app_system_properties' and i != 'tbl_sites' and i != 'tbl_inventory' and i != 'send_messages'
#                 and i != 'send_variablevaluesmessage']
#     else:
#         return no_update
#     # if ipval != '':
#     #     if button > 0:
#     #
#     #         server = SSHTunnelForwarder(
#     #             (ipval, 22),
#     #             ssh_username='soudani',
#     #             ssh_password="univ484067152",
#     #             remote_bind_address=(ipval, 3306))
#     #
#     #         server.start()
#     #
#     #         try:
#     #             conn = mariadb.connect(
#     #                 user="dashapp",
#     #                 password="dashapp",
#     #                 host=ipval,
#     #                 port=3306,
#     #                 database=db_name
#     #             )
#     #
#     #         except mariadb.Error as e:
#     #             print(f"Error connecting to MariaDB Platform: {e}")
#     #             sys.exit(1)
#     #         # Get Cursor
#     #         cur = conn.cursor()
#     #         # cur.execute("SELECT * FROM received_variablevalues WHERE LOCAL_TIMESTAMP <'2020-07-22 18:11:24'")
#     #         b = "select table_name from information_schema.tables where TABLE_SCHEMA='{}'".format(db_name)
#     #         # a = "SELECT DISTINCT VARIABLE_NAME FROM received_variablevalues "
#     #
#     #         cur.execute(b)
#     #         t = cur.fetchall()
#     #         df = pd.DataFrame(t)
#     #         print(df)
#     #         m = []
#     #         for i in t:
#     #             m.append(i[0])
#     #
#     #         return [{'label': i, 'value': i} for i in m if
#     #                 i != 'app_variablerequest' and i != 'send_controlvalues' and i != 'received_ack' and i != 'send_vw_variablerequestdestination' and i != 'flyway_schema_history'
#     #                 and i != 'app_vw_messaging_followup' and i != 'received_variablerequest' and i != 'received_controlvalues' and i != 'app_system_properties'
#     #                 and i != 'tbl_sites' and i != 'tbl_inventory' and i != 'send_messages' and i != 'send_variablevaluesmessage']
#     # else: return no_update
#
# @app.callback(Output('hiddendb1', 'children'),
#               [Input('dbvalchoosen', 'value')], )
# def dbname(dbch):
#     if dbch != None :
#         server = SSHTunnelForwarder(
#             ("193.54.2.211", 22),
#             ssh_username='soudani',
#             ssh_password="univ484067152",
#             remote_bind_address=("193.54.2.211", 3306))
#
#         server.start()
#
#         try:
#             conn = mariadb.connect(
#                 user="dashapp",
#                 password="dashapp",
#                 host="193.54.2.211",
#                 port=3306,
#                 database="rcckn"
#             )
#
#         except mariadb.Error as e:
#             print(f"Error connecting to MariaDB Platform: {e}")
#             sys.exit(1)
#         # Get Cursor
#         cur = conn.cursor()
#         # cur.execute("SELECT * FROM received_variablevalues WHERE LOCAL_TIMESTAMP <'2020-07-22 18:11:24'")
#         # b = "SELECT column_name FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_NAME = '{}' ORDER BY ORDINAL_POSITION".format(
#         #     'received_variablevalues')
#         print("dbch",dbch)
#
#         if dbch == 'send_variablevalues':
#             cur.execute("SELECT * FROM send_variablevalues ")
#             t = cur.fetchall()
#             df = pd.DataFrame(t)
#             df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
#                           'PROCESSED', 'TIMED_OUT','UNREFERENCED']
#             df.to_csv('lermab.csv')
#         if dbch == 'received_variablevalues':
#             cur.execute("SELECT * FROM received_variablevalues ")
#             t = cur.fetchall()
#             df = pd.DataFrame(t)
#             df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
#                           'REMOTE_ID', 'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
#                           'CONVERTED_NUM_VALUE']
#             df.to_csv('lermab.csv')
#
# @app.callback(Output('memory-output', 'data'),
#               Input('dbvalname', 'value'))
# def filter_values(val_selected):
#     if val_selected == None :
#         raise PreventUpdate
#     df = pd.read_csv('lermab.csv')
#     print('val_selected', val_selected)
#     if not val_selected:
#         # Return all the rows on initial load/no country selected.
#
#         return df.to_dict('records')
#
#     filtered = df.query('VARIABLE_NAME in @val_selected').to_dict('records')
#     print('filtered', filtered[:10])
#
#     return filtered
#
# @app.callback([Output('getdbtable', 'data'),
#                Output('getdbtable', 'columns'),],
#               [Input('memory-output', 'data'),
#                Input('dbvalchoosen', 'value'),
#                Input('dbvaldate', 'value')],
#               )
#
# def on_data_set_table(data,dbch,valdate):
#     if data is None or valdate == None:
#         raise PreventUpdate
#     df = pd.DataFrame(data)
#     if dbch == 'received_variablevalues':
#         df['REMOTE_TIMESTAMP'] = df['REMOTE_TIMESTAMP'].astype('string')
#         a = []
#         for col in df['REMOTE_TIMESTAMP'] :
#             a.append(col[:10])
#         df['dates'] = a
#         valdate_new = []
#         for i in range(len(valdate)) :
#             valdate_new.append(valdate[i][:10])
#         df1 = df[df['dates'].isin(valdate_new)]
#         x = df1.to_dict('record')
#         return x, [{'name': i, 'id': i} for i in df.columns if i.startswith('Unn') != 1 or i != 'dates']
#
#     if dbch == 'send_variablevalues':
#         df['TIMESTAMP'] = df['TIMESTAMP'].astype('string')
#         a = []
#         for col in df['TIMESTAMP'] :
#             a.append(col[:10])
#         df['dates'] = a
#         valdate_new = []
#         for i in range(len(valdate)) :
#             valdate_new.append(valdate[i][:10])
#         df1 = df[df['dates'].isin(valdate_new)]
#         x = df1.to_dict('record')
#         return x, [{'name': i, 'id': i} for i in df.columns if i.startswith('Unn') != 1 or i != 'dates']
#
# @app.callback(Output('dbvalname', 'options'),
#               [Input('dbvalchoosen', 'value')], )
# def dbname(dbch):
#     if dbch != None :
#         server = SSHTunnelForwarder(
#             ("193.54.2.211", 22),
#             ssh_username='soudani',
#             ssh_password="univ484067152",
#             remote_bind_address=("193.54.2.211", 3306))
#
#         server.start()
#
#         try:
#             conn = mariadb.connect(
#                 user="dashapp",
#                 password="dashapp",
#                 host="193.54.2.211",
#                 port=3306,
#                 database="rcckn"
#             )
#
#         except mariadb.Error as e:
#             print(f"Error connecting to MariaDB Platform: {e}")
#             sys.exit(1)
#         # Get Cursor
#         cur = conn.cursor()
#         # cur.execute("SELECT * FROM received_variablevalues WHERE LOCAL_TIMESTAMP <'2020-07-22 18:11:24'")
#         # b = "SELECT column_name FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_NAME = '{}' ORDER BY ORDINAL_POSITION".format(
#         #     'received_variablevalues')
#
#         cur.execute("SELECT DISTINCT VARIABLE_NAME FROM {} ".format(dbch))
#         t = cur.fetchall()
#         m = []
#         for i in t:
#             m.append(i[0]) # all variable as tuple, got name with [0]
#         return [{'label': i, 'value': i} for i in m]
#     else:
#         raise PreventUpdate
#
# @app.callback(Output('hiddendb2', 'children'),
#               [Input('memory-output', 'data'),
#                Input('dbvalchoosen', 'value')] )
#
# def vv(data, dbch):
#     if data == [] or data == None:
#         raise PreventUpdate
#     df = pd.DataFrame(data)
#     print('dbch', dbch)
#     if dbch == 'received_variablevalues':
#         # df.columns = ['ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP',
#         #           'REMOTE_ID', 'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT',
#         #           'CONVERTED_NUM_VALUE']
#         df['REMOTE_TIMESTAMP'] = df.REMOTE_TIMESTAMP.apply(pd.to_datetime)
#         df["day"] = df.REMOTE_TIMESTAMP.dt.day
#         df["month"] = df.REMOTE_TIMESTAMP.dt.month
#         df["year"] = df.REMOTE_TIMESTAMP.dt.year
#         a = [str(i) + '-' + str(j) + '-' + str(k) for i, j, k in zip(df["year"], df["month"], df["day"])]
#         a = list(set(a))
#         b = pd.to_datetime(a)
#         b = sorted(b)
#         print('bbbbbbbbbbbb', b)
#         return b
#     elif dbch == "send_variablevalues" :
#         print('buraya giriyor mu')
#         df.TIMESTAMP = df.TIMESTAMP.apply(pd.to_datetime)
#         df["day"] = df.TIMESTAMP.dt.day
#         df["month"] = df.TIMESTAMP.dt.month
#         df["year"] = df.TIMESTAMP.dt.year
#         a = [str(i) + '-' + str(j) + '-' + str(k) for i, j, k in zip(df["year"], df["month"], df["day"])]
#         a = list(set(a))
#         b = pd.to_datetime(a)
#         b = sorted(b)
#         print('bbbbbbbbbbbb22222', b)
#         return b
#
#
# @app.callback(Output('dbvaldate', 'options'),
#               [Input('hiddendb2', 'children')])
# def xx(f):
#     if f == [] or f == None:
#         raise PreventUpdate
#     else:
#         return [{'label': i[:10], 'value': i} for i in f]
#
# @app.callback(Output('getdbgraph', 'figure'),
#               [Input('memory-output', 'data'),
#                Input('dbvalname', 'value'),
#                Input('dbvaldate', 'value'),
#                Input('sliderWidthdb', 'value'),
#                Input('sliderHeightdb', 'value'),],
#               [State('dbvalchoosen', 'value')] )
# def on_data_set_graph(data, valy, valdat,sliderw, sliderh, dbch):
#     if data is None or valy == [] or valdat == [] or valdat == None :
#         raise PreventUpdate
#     df = pd.DataFrame(data)
#     fig = go.Figure()
#     if dbch == 'received_variablevalues':
#
#         df.columns = ['','ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'LOCAL_TIMESTAMP', 'REMOTE_ID',
#                       'REMOTE_TIMESTAMP', 'REMOTE_MESSAGE_ID', 'PROCESSED', 'TIMED_OUT', 'CONVERTED_NUM_VALUE']
#         a = []
#         for col in df['REMOTE_TIMESTAMP']:
#             a.append(col[:10])
#         df['dates'] = a
#         valdate_new = []
#         for i in range(len(valdat)):
#             valdate_new.append(valdat[i][:10])
#         for j in range(len(valy)):
#             for k in range(len(valdate_new)):
#                 a = df[df['VARIABLE_NAME'] == valy[j]]['VARIABLE_NUM_VALUE']
#                 b = df[df['dates'] == valdate_new[k]]['REMOTE_TIMESTAMP']
#                 fig.add_trace(go.Scatter(x=b, y=a, mode='markers', name="{}/{}".format(valy[j], valdate_new[k])))
#             fig.update_layout(
#                 autosize=True,
#                 width=sliderw,
#                 height=sliderh,
#                 margin=dict(
#                     l=50,
#                     r=50,
#                     b=50,
#                     t=50,
#                     pad=4
#                 ),
#                 hovermode='x unified',
#                 uirevision=valy[j], ),
#         return fig
#     else :
#         df.columns = ['','ID', 'VARIABLE_NAME', 'VARIABLE_NUM_VALUE', 'VARIABLE_STR_VALUE', 'TIMESTAMP',
#                       'PROCESSED', 'TIMED_OUT', 'UNREFERENCED']
#         a = []
#         for col in df['TIMESTAMP'] :
#             a.append(col[:10])
#         df['dates'] = a
#         valdate_new = []
#         for i in range(len(valdat)) :
#             valdate_new.append(valdat[i][:10])
#         for j in range(len(valy)):
#             for k in range(len(valdate_new)):
#                 a = df[df['VARIABLE_NAME'] == valy[j]]['VARIABLE_NUM_VALUE']
#                 b = df[df['dates']== valdate_new[k]]['TIMESTAMP']
#                 fig.add_trace(go.Scatter(x=b, y=a, mode='markers', name="{}/{}".format(valy[j],valdate_new[k] )))
#                 fig.update_layout(
#                     autosize=True,
#                     width=sliderw,
#                     height=sliderh,
#                     margin=dict(
#                         l=50,
#                         r=50,
#                         b=50,
#                         t=50,
#                         pad=4
#                     ),
#                     hovermode='x unified',
#                     uirevision=valy[j], ),
#         return fig

if __name__ == '__main__':
    # app.run_server(debug = True)
    app.run_server(debug=True, host='0.0.0.0', port=8049)
